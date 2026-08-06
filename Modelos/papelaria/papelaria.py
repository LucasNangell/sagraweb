import os
import re
import sys
import time
import shutil
import pythoncom
import win32com.client
from openpyxl import load_workbook


ARQUIVO_PLANILHA = "papelaria_variacoes_unicas.xlsx"
ABA_PLANILHA = "Variacoes"


# Ordem amigável para perguntar os dados ao usuário
ORDEM_VARIAVEIS = [
    "nome_deputado",
    "genero",
    "partido",
    "uf",
    "anexo",
    "gabinete",
    "email",
    "nome_assessor",
    "contato_assessor",
    "funcao_assessor",
]


ROTULOS_VARIAVEIS = {
    "nome_deputado": "Nome do deputado/deputada",
    "genero": "Gênero/Cargo. Ex: Deputado Federal ou Deputada Federal",
    "partido": "Partido. Ex: PT, PL, MDB",
    "uf": "UF. Ex: DF, MG, SP",
    "anexo": "Anexo. Ex: IV",
    "gabinete": "Gabinete. Ex: 123",
    "email": "E-mail",
    "nome_assessor": "Nome do assessor/assessora",
    "contato_assessor": "Contato do assessor/assessora",
    "funcao_assessor": "Função do assessor/assessora",
}


def pasta_script():
    return os.path.dirname(os.path.abspath(__file__))


def limpar_nome_arquivo(texto):
    texto = str(texto).strip()
    texto = re.sub(r'[\\/*?:"<>|]', "", texto)
    texto = re.sub(r"\s+", "_", texto)
    return texto


def normalizar_gabinete(valor):
    """
    Aceita:
    - 123
    - Gabinete 123
    - Anexo IV, Gabinete 123

    Retorna sempre os últimos 3 dígitos, com zeros à esquerda se necessário.
    """
    numeros = re.findall(r"\d+", str(valor))

    if not numeros:
        return str(valor).strip()

    numero = numeros[-1]

    if len(numero) > 3:
        numero = numero[-3:]

    return numero.zfill(3)


def normalizar_quebras_para_corel(texto):
    """
    Preserva quebras de linha no CorelDRAW.

    O CorelDRAW normalmente entende melhor chr(13), equivalente a Chr(13) no VBA.
    """
    if texto is None:
        return ""

    texto = str(texto)
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    texto = texto.replace("\n", chr(13))
    return texto


def substituir_variaveis(texto_modelo, dados):
    """
    Substitui variáveis no formato {variavel}.
    Preserva quebras de linha.
    """
    if texto_modelo is None:
        return None

    texto = str(texto_modelo)

    def trocar(match):
        nome_var = match.group(1).strip()

        # Correção tolerante para possível erro de digitação na planilha
        if nome_var == "nome_acessor":
            nome_var = "nome_assessor"

        return str(dados.get(nome_var, ""))

    texto = re.sub(r"\{([^{}]+)\}", trocar, texto)
    return normalizar_quebras_para_corel(texto)


def extrair_variaveis(texto):
    if not texto:
        return set()

    variaveis = set(re.findall(r"\{([^{}]+)\}", str(texto)))

    # Correção tolerante para possível erro de digitação
    if "nome_acessor" in variaveis:
        variaveis.remove("nome_acessor")
        variaveis.add("nome_assessor")

    return variaveis


def dividir_frente_verso(texto):
    """
    Detecta textos no formato:

    Frente:
    texto da frente

    Verso:
    texto do verso

    Retorna:
    {
        1: "texto da frente",
        2: "texto do verso"
    }

    Se não houver Frente/Verso, retorna None.
    """
    if not texto:
        return None

    texto = str(texto)

    # Remove aspas externas caso venham da planilha como parte do texto
    texto = texto.strip()
    if texto.startswith('"') and texto.endswith('"'):
        texto = texto[1:-1]

    texto = texto.replace("\r\n", "\n").replace("\r", "\n")

    padrao = re.compile(
        r"^\s*Frente:\s*\n?(.*?)\n\s*Verso:\s*\n?(.*?)\s*$",
        re.IGNORECASE | re.DOTALL
    )

    match = padrao.match(texto)

    if not match:
        return None

    frente = match.group(1).strip("\n")
    verso = match.group(2).strip("\n")

    return {
        1: frente,
        2: verso,
    }

def encontrar_shapes_por_nome(shapes, nome_procurado, encontrados=None):
    """
    Procura todos os objetos com determinado nome dentro de uma coleção de shapes,
    inclusive dentro de grupos.
    """
    if encontrados is None:
        encontrados = []

    for i in range(1, shapes.Count + 1):
        shape = shapes.Item(i)

        try:
            nome_shape = str(shape.Name).strip()
        except Exception:
            nome_shape = ""

        if nome_shape == nome_procurado:
            encontrados.append(shape)

        # Procura dentro de grupos
        try:
            if shape.Shapes.Count > 0:
                encontrar_shapes_por_nome(shape.Shapes, nome_procurado, encontrados)
        except Exception:
            pass

    return encontrados


def encontrar_shapes_na_pagina(page, nome_procurado):
    """
    Procura todos os objetos com determinado nome em uma página específica.
    """
    encontrados = []

    for l in range(1, page.Layers.Count + 1):
        layer = page.Layers.Item(l)
        encontrar_shapes_por_nome(layer.Shapes, nome_procurado, encontrados)

    return encontrados


def encontrar_shapes_no_documento(doc, nome_procurado):
    """
    Procura todos os objetos com determinado nome em todas as páginas do documento.

    Retorna uma lista de dicionários:
    [
        {
            "pagina": 1,
            "shape": shape
        },
        ...
    ]
    """
    encontrados = []

    for p in range(1, doc.Pages.Count + 1):
        page = doc.Pages.Item(p)
        shapes = encontrar_shapes_na_pagina(page, nome_procurado)

        for shape in shapes:
            encontrados.append({
                "pagina": p,
                "shape": shape
            })

    return encontrados

def carregar_modelos_da_planilha():
    caminho = os.path.join(pasta_script(), ARQUIVO_PLANILHA)

    if not os.path.exists(caminho):
        raise FileNotFoundError(
            f"Planilha não encontrada: {caminho}\n"
            f"Coloque {ARQUIVO_PLANILHA} na mesma pasta do script."
        )

    wb = load_workbook(caminho, data_only=True)
    ws = wb[ABA_PLANILHA]

    # Na sua planilha:
    # Linha 1: cabeçalhos gerais
    # Linha 2: nomes dos campos Corel
    linha_campos = 2

    col_produto = 1
    col_frente_verso = 2
    col_cores = 3
    col_arquivo_modelo = 4
    col_inicio_campos = 5

    modelos = []

    for row in range(3, ws.max_row + 1):
        produto = ws.cell(row=row, column=col_produto).value
        frente_verso = ws.cell(row=row, column=col_frente_verso).value
        cores = ws.cell(row=row, column=col_cores).value
        arquivo_modelo = ws.cell(row=row, column=col_arquivo_modelo).value

        if not arquivo_modelo:
            continue

        arquivo_modelo = str(arquivo_modelo).strip()

        # Se algum modelo estiver sem .cdr na planilha, adiciona automaticamente
        if not arquivo_modelo.lower().endswith(".cdr"):
            arquivo_modelo += ".cdr"

        campos = {}

        for col in range(col_inicio_campos, ws.max_column + 1):
            nome_campo_corel = ws.cell(row=linha_campos, column=col).value
            valor_modelo = ws.cell(row=row, column=col).value

            if not nome_campo_corel or valor_modelo is None or str(valor_modelo).strip() == "":
                continue

            campos[str(nome_campo_corel).strip()] = str(valor_modelo)

        if not campos:
            continue

        modelos.append({
            "linha": row,
            "produto": str(produto or "").strip(),
            "frente_verso": str(frente_verso or "").strip(),
            "cores": str(cores or "").strip(),
            "arquivo_modelo": arquivo_modelo,
            "campos": campos,
        })

    return modelos


def listar_modelos(modelos):
    print("\n=== Modelos disponíveis ===\n")

    for i, modelo in enumerate(modelos, start=1):
        print(
            f"{i:02d} - {modelo['produto']} | "
            f"{modelo['frente_verso']} | "
            f"{modelo['cores']} | "
            f"{modelo['arquivo_modelo']}"
        )


def escolher_modelo(modelos):
    listar_modelos(modelos)

    while True:
        escolha = input("\nDigite o número do modelo desejado: ").strip()

        if not escolha.isdigit():
            print("Digite apenas o número do modelo.")
            continue

        indice = int(escolha)

        if indice < 1 or indice > len(modelos):
            print("Modelo inválido.")
            continue

        return modelos[indice - 1]


def variaveis_necessarias(modelo):
    variaveis = set()

    for texto in modelo["campos"].values():
        variaveis.update(extrair_variaveis(texto))

    return variaveis


def perguntar_dados(variaveis):
    dados = {}

    print("\n=== Dados necessários ===\n")

    variaveis_ordenadas = []

    for var in ORDEM_VARIAVEIS:
        if var in variaveis:
            variaveis_ordenadas.append(var)

    for var in sorted(variaveis):
        if var not in variaveis_ordenadas:
            variaveis_ordenadas.append(var)

    for var in variaveis_ordenadas:
        rotulo = ROTULOS_VARIAVEIS.get(var, var)
        valor = input(f"{rotulo}: ").strip()

        if var == "gabinete":
            valor = normalizar_gabinete(valor)

        if var == "uf":
            valor = valor.upper()

        dados[var] = valor

    return dados


def encontrar_shape_por_nome(shapes, nome_procurado):
    """
    Procura uma shape pelo nome, inclusive dentro de grupos.
    """
    for i in range(1, shapes.Count + 1):
        shape = shapes.Item(i)

        try:
            if str(shape.Name).strip() == nome_procurado:
                return shape
        except Exception:
            pass

        try:
            if shape.Shapes.Count > 0:
                encontrada = encontrar_shape_por_nome(shape.Shapes, nome_procurado)
                if encontrada:
                    return encontrada
        except Exception:
            pass

    return None


def encontrar_shape_na_pagina(page, nome_procurado):
    """
    Procura uma shape em todas as camadas de uma página específica.
    """
    for l in range(1, page.Layers.Count + 1):
        layer = page.Layers.Item(l)
        shape = encontrar_shape_por_nome(layer.Shapes, nome_procurado)

        if shape:
            return shape

    return None


def encontrar_shape_no_documento(doc, nome_procurado):
    """
    Procura uma shape em todas as páginas e camadas do documento.
    """
    for p in range(1, doc.Pages.Count + 1):
        page = doc.Pages.Item(p)
        shape = encontrar_shape_na_pagina(page, nome_procurado)

        if shape:
            return shape

    return None


def alterar_texto(shape, novo_texto):
    """
    Altera apenas o texto do objeto, preservando a formatação geral da caixa.

    Para textos com múltiplas linhas, o texto já deve vir com chr(13).
    """
    texto = normalizar_quebras_para_corel(novo_texto)

    try:
        shape.Text.Story.Text = texto
        return
    except Exception as erro_1:
        print(f"Aviso: falha usando shape.Text.Story.Text em '{shape.Name}': {erro_1}")

    try:
        shape.Text.Story.Characters.All.Text = texto
        return
    except Exception as erro_2:
        print(f"Aviso: falha usando Characters.All.Text em '{shape.Name}': {erro_2}")

    try:
        shape.Text.Story = texto
        return
    except Exception as erro_3:
        print(f"Aviso: falha usando shape.Text.Story em '{shape.Name}': {erro_3}")

    raise RuntimeError(
        f"Não foi possível alterar o texto do objeto '{shape.Name}'. "
        "Verifique se ele é um objeto de texto editável e não está convertido em curvas."
    )


def aplicar_campo(doc, nome_campo_corel, texto_modelo, dados):
    """
    Aplica o texto de um campo do Corel.

    Regra:
    - Procura todos os objetos com o nome configurado na planilha.
    - Se a célula tiver Frente:/Verso:, aplica:
        página 1 = Frente
        página 2 = Verso
    - Se não tiver Frente:/Verso:, aplica o mesmo texto em todos os objetos encontrados.
    """

    ocorrencias = encontrar_shapes_no_documento(doc, nome_campo_corel)

    if not ocorrencias:
        raise RuntimeError(
            f"Objeto '{nome_campo_corel}' não encontrado em nenhuma página do documento."
        )

    frente_verso = dividir_frente_verso(texto_modelo)

    if frente_verso:
        total_aplicado = 0

        for ocorrencia in ocorrencias:
            numero_pagina = ocorrencia["pagina"]
            shape = ocorrencia["shape"]

            if numero_pagina == 1:
                texto_final = substituir_variaveis(frente_verso[1], dados)
            elif numero_pagina == 2:
                texto_final = substituir_variaveis(frente_verso[2], dados)
            else:
                # Para páginas extras, mantém a lógica:
                # se houver mais páginas, usa o texto da frente por segurança.
                texto_final = substituir_variaveis(frente_verso[1], dados)

            alterar_texto(shape, texto_final)
            total_aplicado += 1

        print(
            f"Campo '{nome_campo_corel}' aplicado em {total_aplicado} ocorrência(s) "
            f"com regra Frente/Verso."
        )

        return

    texto_final = substituir_variaveis(texto_modelo, dados)

    for ocorrencia in ocorrencias:
        alterar_texto(ocorrencia["shape"], texto_final)

    print(
        f"Campo '{nome_campo_corel}' aplicado em {len(ocorrencias)} ocorrência(s) "
        f"em todas as páginas."
    )


def configurar_pdf(corel, doc):
    """
    Configura PDF para exportar o documento inteiro, incluindo frente e verso.
    Usa as constantes nativas do CorelDRAW quando disponíveis.
    """
    try:
        pdf = doc.PDFSettings
    except Exception:
        try:
            pdf = doc.pdfSettings
        except Exception:
            return

    try:
        constantes = win32com.client.constants
    except Exception:
        constantes = None

    configuracoes_simples = {
        "Author": os.environ.get("USERNAME", "Python"),
        "Subject": "Exportação automática",
        "Keywords": "PDF, CorelDRAW, Macro",
        "JPEGQualityFactor": 90,
        "EmbedFonts": True,
        "SubsetFonts": True,
        "TextAsCurves": False,
    }

    for atributo, valor in configuracoes_simples.items():
        try:
            setattr(pdf, atributo, valor)
        except Exception:
            pass

    # Equivalente ao VBA:
    # .PublishRange = pdfWholeDocument
    try:
        pdf.PublishRange = constantes.pdfWholeDocument
    except Exception:
        try:
            # Fallback comum em algumas versões do CorelDRAW
            pdf.PublishRange = 0
        except Exception:
            pass

    # Equivalente ao VBA:
    # .BitmapCompression = pdfJPEG
    try:
        pdf.BitmapCompression = constantes.pdfJPEG
    except Exception:
        pass

    # Equivalente ao VBA:
    # .ColorMode = pdfRGB
    try:
        pdf.ColorMode = constantes.pdfRGB
    except Exception:
        pass

    # Garante que todas as páginas estejam dentro do intervalo de publicação.
    # Algumas versões respeitam melhor PageRange quando PublishRange não pega.
    try:
        total_paginas = doc.Pages.Count
        pdf.PageRange = f"1-{total_paginas}"
    except Exception:
        pass


def salvar_documento(doc):
    """
    Salva o documento já aberto.

    O script trabalha sempre em cima de uma cópia do modelo.
    Assim evitamos usar SaveAs, que deu erro via COM no seu Corel.
    """
    try:
        doc.Save()
    except Exception as erro:
        raise RuntimeError(f"Não foi possível salvar o CDR com doc.Save(): {erro}")


def exportar_pdf(corel, doc, caminho_pdf):
    caminho_pdf = os.path.abspath(caminho_pdf)

    configurar_pdf(corel, doc)

    if os.path.exists(caminho_pdf):
        try:
            os.remove(caminho_pdf)
        except Exception:
            pass

    try:
        doc.PublishToPDF(caminho_pdf)
    except Exception as erro:
        raise RuntimeError(f"Não foi possível exportar o PDF com PublishToPDF(): {erro}")

def nome_saida(modelo, dados):
    produto = limpar_nome_arquivo(modelo["produto"] or "Papelaria")
    cores = limpar_nome_arquivo(modelo["cores"] or "")
    nome = limpar_nome_arquivo(dados.get("nome_deputado", "sem_nome"))
    gabinete = limpar_nome_arquivo(dados.get("gabinete", ""))

    partes = [produto, nome]

    if gabinete:
        partes.append(f"Gab{gabinete}")

    if cores:
        partes.append(cores)

    return "_".join(partes)


def gerar_arquivo(modelo, dados):
    base = pasta_script()

    caminho_modelo = os.path.join(base, modelo["arquivo_modelo"])

    if not os.path.exists(caminho_modelo):
        raise FileNotFoundError(
            f"Arquivo modelo não encontrado: {caminho_modelo}\n"
            "Verifique se o arquivo .cdr está na mesma pasta do script."
        )

    base_saida = nome_saida(modelo, dados)

    caminho_saida_cdr = os.path.join(base, f"{base_saida}.cdr")
    caminho_saida_pdf = os.path.join(base, f"{base_saida}.pdf")

    if os.path.exists(caminho_saida_cdr):
        os.remove(caminho_saida_cdr)

    shutil.copy2(caminho_modelo, caminho_saida_cdr)

    pythoncom.CoInitialize()

    doc = None

    try:
        corel = win32com.client.gencache.EnsureDispatch("CorelDRAW.Application")
        corel.Visible = True

        doc = corel.OpenDocument(os.path.abspath(caminho_saida_cdr))
        print(f"\nTotal de páginas do modelo aberto: {doc.Pages.Count}")

        print("\n=== Preenchendo campos no CorelDRAW ===\n")

        for nome_campo_corel, texto_modelo in modelo["campos"].items():
            print(f"Aplicando campo: {nome_campo_corel}")
            aplicar_campo(doc, nome_campo_corel, texto_modelo, dados)

        time.sleep(0.5)

        salvar_documento(doc)

        time.sleep(0.5)

        exportar_pdf(corel, doc, caminho_saida_pdf)

        print("\nArquivo gerado com sucesso!")
        print(f"CDR: {caminho_saida_cdr}")
        print(f"PDF: {caminho_saida_pdf}")

    finally:
        if doc is not None:
            try:
                doc.Close()
            except Exception:
                pass

        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def main():
    print("=== Gerador Universal de Papelaria CorelDRAW ===")

    modelos = carregar_modelos_da_planilha()

    if not modelos:
        raise RuntimeError("Nenhum modelo válido foi encontrado na planilha.")

    modelo = escolher_modelo(modelos)

    print("\n=== Modelo selecionado ===")
    print(f"Produto: {modelo['produto']}")
    print(f"Frente/Verso: {modelo['frente_verso']}")
    print(f"Cores: {modelo['cores']}")
    print(f"Arquivo modelo: {modelo['arquivo_modelo']}")

    variaveis = variaveis_necessarias(modelo)
    dados = perguntar_dados(variaveis)

    # Regra útil:
    # se o modelo usa {anexo}, mas o usuário não informou,
    # assume IV como padrão da Câmara.
    if "anexo" in variaveis and not dados.get("anexo"):
        dados["anexo"] = "IV"

    gerar_arquivo(modelo, dados)


if __name__ == "__main__":
    try:
        main()
    except Exception as erro:
        print("\nERRO:")
        print(erro)
        sys.exit(1)