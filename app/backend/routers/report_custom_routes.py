from fastapi import APIRouter, HTTPException, Body
from fastapi.responses import HTMLResponse, StreamingResponse
from typing import List, Optional, Any, Dict, Union
from pydantic import BaseModel
from database import db
import logging
import datetime
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter()
logger = logging.getLogger(__name__)

class ReportColumn(BaseModel):
    id: str
    label: str
    db_field: str

class ReportRequest(BaseModel):
    filters: Dict[str, Any]
    columns: List[ReportColumn]
    generated_by: Optional[str] = None
    sort_by: Optional[str] = None
    sort_order: Optional[str] = "DESC"

def fetch_report_data(req: ReportRequest) -> List[dict]:
    # 1. Build Query
    select_parts = []
    for col in req.columns:
        select_parts.append(f"{col.db_field} AS `{col.label}`")
    
    base_query = """
        FROM tabProtocolos AS p
        LEFT JOIN tabDetalhesServico AS d ON (p.NroProtocolo = d.NroProtocoloLinkDet) AND (p.AnoProtocolo = d.AnoProtocoloLinkDet)
        LEFT JOIN tabAndamento AS a ON (p.NroProtocolo = a.NroProtocoloLink) AND (p.AnoProtocolo = a.AnoProtocoloLink)
        LEFT JOIN tabMidiaDigital AS md ON (p.NroProtocolo = md.NroProtocoloLinkMidia) AND (p.AnoProtocolo = md.AnoProtocoloLinkMidia)
    """
    
    filters = " WHERE a.UltimoStatus = 1"
    params = {}
    
    # Apply Filters (Same logic as search_os)
    f = req.filters
    if isinstance(f.get('include_finished'), str):
            include_finished = f.get('include_finished').lower() == 'true'
    else:
            include_finished = bool(f.get('include_finished'))
            
    if not include_finished:
        filters += " AND a.SituacaoLink NOT IN ('Entregue', 'Cancelada', 'Cancelado')"

    if f.get('nr_os'):
        filters += " AND p.NroProtocolo = %(nr_os)s"
        params['nr_os'] = f.get('nr_os')
    if f.get('ano'):
        filters += " AND p.AnoProtocolo = %(ano)s"
        params['ano'] = f.get('ano')
    if f.get('produto'):
        filters += " AND d.TipoPublicacaoLink LIKE %(produto)s"
        params['produto'] = f"%{f.get('produto')}%"
    if f.get('titulo'):
        filters += " AND d.Titulo LIKE %(titulo)s"
        params['titulo'] = f"%{f.get('titulo')}%"
    if f.get('solicitante'):
        filters += " AND p.NomeUsuario LIKE %(solicitante)s"
        params['solicitante'] = f"%{f.get('solicitante')}%"

    # List filters
    situacao = f.get('situacao')
    if situacao:
        if isinstance(situacao, str): situacao = [situacao]
        if len(situacao) > 0:
            clean_list = [s for s in situacao if s] # Basic cleanup
            if clean_list:
                placeholders = ', '.join([f'%(situacao_{i})s' for i in range(len(clean_list))])
                filters += f" AND a.SituacaoLink IN ({placeholders})"
                for i, s in enumerate(clean_list):
                    params[f'situacao_{i}'] = s

    setor = f.get('setor')
    if setor:
        if isinstance(setor, str): setor = [setor]
        if len(setor) > 0:
            clean_list = [s for s in setor if s]
            if clean_list:
                placeholders = ', '.join([f'%(setor_{i})s' for i in range(len(clean_list))])
                filters += f" AND a.SetorLink IN ({placeholders})"
                for i, s in enumerate(clean_list):
                    params[f'setor_{i}'] = s
    
    prioridade = f.get('prioridade')
    if prioridade:
        if isinstance(prioridade, str): prioridade = [prioridade]
        if len(prioridade) > 0:
            clean_list = [s for s in prioridade if s]
            if clean_list:
                placeholders = ', '.join([f'%(prio_{i})s' for i in range(len(clean_list))])
                filters += f" AND p.EntregPrazoLink IN ({placeholders})"
                for i, s in enumerate(clean_list):
                    params[f'prio_{i}'] = s
    
    # Ordering
    sort_by = req.sort_by
    sort_order = req.sort_order if req.sort_order in ['ASC', 'DESC'] else 'DESC'
    
    if sort_by and sort_by.strip():
         # Basic sanitization to prevent injection if it's not a known field
         # Ideally we validate against a whitelist of columns, but for now specific characters check
         # or rely on frontend sending valid "table.column"
         safe_sort = sort_by.replace(';', '').replace('DROP', '').replace('DELETE', '').strip()
         
         # Special case: default composite sort
         if 'Num' in safe_sort or 'Ano' in safe_sort:
             order_by = f" ORDER BY {safe_sort} {sort_order}"
         else:
             order_by = f" ORDER BY {safe_sort} {sort_order}, p.AnoProtocolo DESC, p.NroProtocolo DESC"
    else:
        order_by = " ORDER BY p.AnoProtocolo DESC, p.NroProtocolo DESC, p.DataEntrada DESC"

    full_query = f"SELECT {', '.join(select_parts)} {base_query} {filters} {order_by}"
    
    return db.execute_query(full_query, params)


@router.post("/reports/custom")
def generate_custom_report(req: ReportRequest):
    try:
        rows = fetch_report_data(req)
        return HTMLResponse(content=render_report_html(rows, req.columns, req.generated_by))
    except Exception as e:
        logger.error(f"Error generating HTML report: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/reports/custom/xlsx")
def generate_custom_report_xlsx(req: ReportRequest):
    try:
        rows = fetch_report_data(req)
        
        # Create Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Relatório"
        
        # Headers
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
        
        for col_idx, col in enumerate(req.columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=col.label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            
        # Data
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, col in enumerate(req.columns, 1):
                val = row_data.get(col.label)
                
                # Format dates
                if isinstance(val, (datetime.date, datetime.datetime)):
                     # Excel handles date objects natively if we keep them as objects, 
                     # but string format is safer for consistency unless we set number_format
                     val = val.strftime("%d/%m/%Y")
                     
                ws.cell(row=row_idx, column=col_idx, value=val)
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter # Get the column name
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = min(adjusted_width, 50) # Cap width

        # Save to buffer
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"Relatorio_Sagra_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        
        headers = {
            'Content-Disposition': f'attachment; filename="{filename}"'
        }
        
        return StreamingResponse(output, headers=headers, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    except Exception as e:
        logger.error(f"Error generating XLSX report: {e}")
        raise HTTPException(status_code=500, detail=str(e))


def render_report_html(rows: List[dict], columns: List[ReportColumn], user: str):
    # CSS for printing support
    css = """
    body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; font-size: 12px; margin: 20px; }
    h1 { color: #2c3e50; border-bottom: 2px solid #2ecc71; padding-bottom: 10px; }
    .meta { color: #7f8c8d; margin-bottom: 20px; font-size: 0.9em; }
    table { width: 100%; border-collapse: collapse; margin-top: 10px; }
    th { background-color: #f2f2f2; border: 1px solid #ddd; padding: 8px; text-align: left; font-weight: 600; color: #333; }
    td { border: 1px solid #ddd; padding: 6px 8px; vertical-align: top; }
    tr:nth-child(even) { background-color: #f9f9f9; }
    tr:hover { background-color: #f1f1f1; }
    @media print {
        @page { size: landscape; margin: 10mm; }
        body { margin: 0; -webkit-print-color-adjust: exact; }
    }
    """
    
    headers = "".join([f"<th>{col.label}</th>" for col in columns])
    
    rows_html = ""
    for row in rows:
        cells = ""
        for col in columns:
            val = row.get(col.label)
            
            # Format dates nicely
            if isinstance(val, datetime.date) or isinstance(val, datetime.datetime):
                val = val.strftime("%d/%m/%Y")
            
            # None to empty
            if val is None: val = ""
            
            cells += f"<td>{val}</td>"
        rows_html += f"<tr>{cells}</tr>"
        
    date_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    user_str = f"Gerado por: {user} - " if user else ""
    total = len(rows)

    html = f"""
    <!DOCTYPE html>
    <html lang="pt-BR">
    <head>
        <meta charset="UTF-8">
        <title>Relatório Personalizado - SAGRA</title>
        <style>{css}</style>
    </head>
    <body>
        <h1>Relatório de Protocolos</h1>
        <div class="meta">
            {user_str}{date_str}<br>
            <strong>Total de Registros: {total}</strong>
        </div>
        <table>
            <thead>
                <tr>{headers}</tr>
            </thead>
            <tbody>
                {rows_html}
            </tbody>
        </table>
        <script>
            // Auto open print dialog optionally?
            // window.print();
        </script>
    </body>
    </html>
    """
    return html

