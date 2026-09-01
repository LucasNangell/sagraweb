"""Exportação XLSX da lista de OSs — mesmos filtros de GET /api/os/search."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from services.os_search_service import OS_SEARCH_MAX_LIMIT, execute_os_search

OS_SEARCH_EXPORT_MAX_ROWS = 10_000

OS_LIST_EXPORT_COLUMNS: Tuple[Tuple[str, str], ...] = (
    ("Nr OS", "nr_os"),
    ("Ano", "ano"),
    ("Entrada", "data_entrada"),
    ("Produto", "produto"),
    ("Solicitante", "solicitante"),
    ("Título", "titulo"),
    ("Situação", "situacao"),
    ("Setor", "setor"),
    ("Prioridade", "prioridade"),
    ("Data", "data_entrega"),
)


def fetch_all_os_search_rows(**search_kwargs: Any) -> List[Dict[str, Any]]:
    """Busca todas as OSs filtradas paginando internamente (ignora page/limit do cliente)."""
    all_rows: List[Dict[str, Any]] = []
    page = 1
    while len(all_rows) < OS_SEARCH_EXPORT_MAX_ROWS:
        result = execute_os_search(
            **search_kwargs,
            page=page,
            limit=OS_SEARCH_MAX_LIMIT,
            simple_pagination=True,
            include_total=False,
        )
        batch = result.get("data") or []
        if not batch:
            break
        all_rows.extend(batch)
        if len(all_rows) >= OS_SEARCH_EXPORT_MAX_ROWS:
            all_rows = all_rows[:OS_SEARCH_EXPORT_MAX_ROWS]
            break
        if not result.get("meta", {}).get("has_next"):
            break
        page += 1
    return all_rows


def _format_export_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return value if value is not None else ""


def build_os_list_xlsx(rows: List[Dict[str, Any]]) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lista de OS"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="1E4A7A", end_color="1E4A7A", fill_type="solid"
    )

    for col_idx, (label, _) in enumerate(OS_LIST_EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, (_, key) in enumerate(OS_LIST_EXPORT_COLUMNS, 1):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=_format_export_cell(row.get(key)),
            )

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
