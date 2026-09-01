"""
Modelos Corel de papelaria: leitura da planilha, importação MySQL, matching OS e prefill.
Sem automação COM — geração real fica no worker Windows (fase posterior).
"""
from __future__ import annotations

import json
import logging
import re
import unicodedata
from pathlib import Path
from typing import Any, Optional

from config_manager import PROJECT_ROOT
from database import db

logger = logging.getLogger(__name__)

# --- Caminhos (servidor não expõe UNC ao frontend) ---
SPREADSHEET_PATH = PROJECT_ROOT / "Modelos" / "papelaria" / "papelaria_variacoes_unicas.xlsx"
ABA_PLANILHA = "Variacoes"
COREL_MODELS_DFS_DIR = (
    r"\\redecamara\DfsData\CGraf\Sefoc\Laboratorio\Modelos\Modelos Corel\Papelaria\CSP\WEB"
)
COREL_MODELS_DEV_DIR = PROJECT_ROOT / "Modelos" / "papelaria" / "cdr"

ORDEM_VARIAVEIS = [
    "nome_deputado",
    "genero",
    "partido",
    "uf",
    "anexo",
    "gabinete",
    "email",
    "nome_assessor",
    "contato_assessor",
    "funcao_assessor",
]

ROTULOS_VARIAVEIS: dict[str, str] = {
    "nome_deputado": "Nome do deputado/deputada",
    "genero": "Gênero/Cargo. Ex: Deputado Federal ou Deputada Federal",
    "partido": "Partido. Ex: PT, PL, MDB",
    "uf": "UF. Ex: DF, MG, SP",
    "anexo": "Anexo. Ex: IV",
    "gabinete": "Gabinete. Ex: 123",
    "email": "E-mail",
    "nome_assessor": "Nome do assessor/assessora",
    "contato_assessor": "Contato do assessor/assessora",
    "funcao_assessor": "Função do assessor/assessora",
}

# Templates canônicos por nome de objeto Corel (planilha / modelos legados).
# Aplicados somente quando o texto inspecionado NÃO contém `{variavel}`.
_COREL_ENDERECO_PT = (
    "Câmara dos Deputados\nAnexo {anexo}, Gabinete {gabinete}\nCEP 70160-900 - Brasília/DF"
)
_COREL_ENDERECO_EN = (
    "Chamber of Deputies\nAnnex {anexo}, Office {gabinete}\nZip Code 70160-900 - Brasília/DF"
)
_COREL_CONTATO_PT = "Fones: (61) 3215-5{gabinete}\n3215-3{gabinete}\n{email}"
_COREL_CONTATO_EN = "Phones: (61) 3215-5{gabinete}\n3215-3{gabinete}\n{email}"
_COREL_RODAPE_PT = (
    "Câmara dos Deputados | Anexo {anexo} - Gabinete {gabinete} | "
    "CEP 70160-900 - Brasília/DF\nTels (61) 3215-5{gabinete}/3{gabinete} | {email}"
)
_COREL_CABECALHO = "Gabinete {artigo} {genero_curto}\n{nome_deputado}\n- {partido}/{uf}"
_CABECALHO_LEGADO = "Gabinete do Deputado"
_CABECALHO_DERIVADO = "Gabinete {artigo} {genero_curto}"
VARIAVEIS_DERIVADAS = frozenset({"artigo", "genero_curto"})

COREL_CAMPO_TEMPLATES_CANONICOS: dict[str, str] = {
    "nomedep": "{nome_deputado}",
    "genero": "{genero}",
    "partidodep": "{partido}/{uf}",
    "nomeassessor": "{nome_assessor}",
    "contatoassessor": "{contato_assessor}",
    "funcao": "{funcao_assessor}",
    "cabecalho": _COREL_CABECALHO,
    "rodape": _COREL_RODAPE_PT,
    "endereco": _COREL_ENDERECO_PT,
    "contato": _COREL_CONTATO_PT,
}

_table_ready = False


# ---------------------------------------------------------------------------
# Utilitários (reaproveitados de Modelos/papelaria/papelaria.py)
# ---------------------------------------------------------------------------


def _as_str(value: Any, default: str = "") -> str:
    """Coerce valor do MySQL (int, etc.) para str antes de .strip()."""
    if value is None:
        return default
    return str(value).strip()


def limpar_nome_arquivo(texto: str) -> str:
    texto = str(texto).strip()
    texto = re.sub(r'[\\/*?:"<>|]', "", texto)
    texto = re.sub(r"\s+", "_", texto)
    return texto


def normalizar_gabinete(valor: str) -> str:
    numeros = re.findall(r"\d+", str(valor))
    if not numeros:
        return str(valor).strip()
    numero = numeros[-1]
    if len(numero) > 3:
        numero = numero[-3:]
    return numero.zfill(3)


def normalizar_quebras_para_corel(texto: Optional[str]) -> str:
    if texto is None:
        return ""
    texto = str(texto)
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    return texto.replace("\n", chr(13))


def enriquecer_variaveis_derivadas(dados: dict[str, Any]) -> dict[str, Any]:
    out = dict(dados)
    feminino = "deputada" in str(out.get("genero") or "").lower()
    out["artigo"] = "da" if feminino else "do"
    out["genero_curto"] = "Deputada" if feminino else "Deputado"
    return out


def extrair_variaveis(texto: Optional[str]) -> set[str]:
    if not texto:
        return set()
    variaveis = set(re.findall(r"\{([^{}]+)\}", str(texto)))
    if "nome_acessor" in variaveis:
        variaveis.remove("nome_acessor")
        variaveis.add("nome_assessor")
    return variaveis - VARIAVEIS_DERIVADAS


def parse_inline_bold(texto: str) -> tuple[str, list[tuple[int, int]]]:
    """Remove pares `**...**` e devolve texto plano + intervalos 0-based em negrito."""
    plain_parts: list[str] = []
    ranges: list[tuple[int, int]] = []
    plain_len = 0
    i = 0
    n = len(texto)
    while i < n:
        if texto.startswith("**", i):
            close = texto.find("**", i + 2)
            if close < 0:
                plain_parts.append(texto[i:])
                break
            inner = texto[i + 2 : close]
            if inner:
                start = plain_len
                plain_parts.append(inner)
                plain_len += len(inner)
                ranges.append((start, start + len(inner)))
            i = close + 2
            continue
        plain_parts.append(texto[i])
        plain_len += 1
        i += 1
    return "".join(plain_parts), ranges


def _substituir_placeholders(texto_modelo: str, dados: dict[str, Any]) -> str:
    resolvidos = enriquecer_variaveis_derivadas(dados)

    def trocar(match: re.Match[str]) -> str:
        nome_var = match.group(1).strip()
        if nome_var == "nome_acessor":
            nome_var = "nome_assessor"
        return str(resolvidos.get(nome_var, ""))

    return re.sub(r"\{([^{}]+)\}", trocar, texto_modelo)


def substituir_variaveis(texto_modelo: Optional[str], dados: dict[str, Any]) -> Optional[str]:
    if texto_modelo is None:
        return None

    texto = _substituir_placeholders(str(texto_modelo), dados)
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    plain, _ranges = parse_inline_bold(texto)
    return normalizar_quebras_para_corel(plain).replace(chr(13), "\n")


def normalize_corel_object_name(nome: str) -> str:
    t = unicodedata.normalize("NFKD", str(nome).strip())
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]", "", t.lower())


def _looks_english_campo_text(texto: str) -> bool:
    t = str(texto or "").lower()
    return bool(
        re.search(
            r"\b(chamber|annex|office|phones?|zip\s*code|federal\s+deputy)\b",
            t,
        )
    )


def resolve_canonical_campo_template(nome_campo: str, texto_inspect: str) -> Optional[str]:
    """
    Retorna template canônico para objetos Corel conhecidos quando o texto
    inspecionado não contém placeholders `{variavel}` (texto renderizado no .cdr).
    Preserva templates explícitos da planilha quando já possuem `{...}`.
    """
    if extrair_variaveis(texto_inspect):
        return None

    key = normalize_corel_object_name(nome_campo)
    if not key:
        return None

    if key == "endereco":
        return _COREL_ENDERECO_EN if _looks_english_campo_text(texto_inspect) else _COREL_ENDERECO_PT
    if key == "contato":
        return _COREL_CONTATO_EN if _looks_english_campo_text(texto_inspect) else _COREL_CONTATO_PT

    return COREL_CAMPO_TEMPLATES_CANONICOS.get(key)


def slugify_variavel_campo(nome: str, used: Optional[set[str]] = None) -> str:
    base = re.sub(r"[^\w\s-]", "", str(nome).strip())
    base = re.sub(r"\s+", "_", base).strip("_")[:64] or "campo"
    if not re.match(r"^[a-zA-Z]", base):
        base = f"c_{base}"
    if used is None:
        return base
    candidate = base
    suffix = 2
    while candidate in used:
        candidate = f"{base}_{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def normalize_campos_from_inspect(
    campos: dict[str, str],
) -> tuple[dict[str, str], list[str], dict[str, str]]:
    """Normaliza campos Corel sem `{variavel}` para placeholders editáveis."""
    normalized: dict[str, str] = {}
    variaveis_set: set[str] = set()
    labels: dict[str, str] = {}
    used_slugs: set[str] = set()

    for nome_campo, texto_raw in sorted(campos.items()):
        texto = str(texto_raw or "").strip()
        vars_in_text = extrair_variaveis(texto)
        if vars_in_text:
            normalized[nome_campo] = texto
            variaveis_set.update(vars_in_text)
            continue

        canonical = resolve_canonical_campo_template(nome_campo, texto)
        if canonical:
            normalized[nome_campo] = canonical
            variaveis_set.update(extrair_variaveis(canonical))
            continue

        slug = slugify_variavel_campo(nome_campo, used_slugs)
        normalized[nome_campo] = f"{{{slug}}}"
        variaveis_set.add(slug)
        labels[slug] = nome_campo

    return normalized, _order_variaveis_list(list(variaveis_set)), labels


def _order_variaveis_list(variaveis: list[str]) -> list[str]:
    ordenadas: list[str] = []
    seen: set[str] = set()
    for var in ORDEM_VARIAVEIS:
        if var in variaveis and var not in seen and var not in VARIAVEIS_DERIVADAS:
            ordenadas.append(var)
            seen.add(var)
    for var in sorted(variaveis):
        if var not in seen and var not in VARIAVEIS_DERIVADAS:
            ordenadas.append(var)
            seen.add(var)
    return ordenadas


def dividir_frente_verso(texto: Optional[str]) -> Optional[dict[int, str]]:
    if not texto:
        return None
    texto = str(texto).strip()
    if texto.startswith('"') and texto.endswith('"'):
        texto = texto[1:-1]
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    padrao = re.compile(
        r"^\s*Frente:\s*\n?(.*?)\n\s*Verso:\s*\n?(.*?)\s*$",
        re.IGNORECASE | re.DOTALL,
    )
    match = padrao.match(texto)
    if not match:
        return None
    return {1: match.group(1).strip("\n"), 2: match.group(2).strip("\n")}


def normalizar_texto_produto(texto: str) -> str:
    if not texto:
        return ""
    t = unicodedata.normalize("NFKD", str(texto))
    t = "".join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^a-zA-Z0-9\s]", " ", t)
    return " ".join(t.lower().split())


def _parse_json_field(raw: Any) -> Any:
    if raw is None:
        return None
    if isinstance(raw, (dict, list)):
        return raw
    if isinstance(raw, str):
        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            return None
    return None


# ---------------------------------------------------------------------------
# Schema / tabela
# ---------------------------------------------------------------------------


def _table_has_column(column: str) -> bool:
    rows = db.execute_query(
        """
        SELECT COUNT(*) AS cnt
        FROM information_schema.COLUMNS
        WHERE TABLE_SCHEMA = DATABASE()
          AND TABLE_NAME = 'tabPapelariaModelos'
          AND COLUMN_NAME = %(col)s
        """,
        {"col": column},
    )
    return bool(rows and rows[0].get("cnt"))


def ensure_papelaria_modelos_table() -> None:
    global _table_ready
    if _table_ready:
        return

    sql_path = PROJECT_ROOT / "deploy" / "sql" / "create_papelaria_modelos.sql"
    if sql_path.is_file():
        raw = sql_path.read_text(encoding="utf-8")
        for stmt in raw.split(";"):
            s = stmt.strip()
            if s:
                db.execute_query(s)

    legacy_migrations: list[tuple[str, str]] = [
        ("Produto", "VARCHAR(120) NULL"),
        ("FrenteVerso", "VARCHAR(32) NULL"),
        ("Cores", "VARCHAR(64) NULL"),
        ("ArquivoModelo", "VARCHAR(255) NULL"),
        ("NomeExibicao", "VARCHAR(255) NULL"),
        ("CamposJson", "JSON NULL"),
        ("VariaveisJson", "JSON NULL"),
        ("MockupConfigJson", "JSON NULL"),
        ("PreviewConfigJson", "JSON NULL"),
        ("Ativo", "TINYINT(1) NOT NULL DEFAULT 1"),
        ("Ordem", "INT NOT NULL DEFAULT 0"),
        ("AtualizadoEm", "TIMESTAMP NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP"),
    ]
    for col, definition in legacy_migrations:
        if not _table_has_column(col):
            try:
                db.execute_query(f"ALTER TABLE tabPapelariaModelos ADD COLUMN {col} {definition}")
            except Exception as exc:
                logger.warning("Migração coluna %s em tabPapelariaModelos: %s", col, exc)

    _migrate_cabecalho_genero_placeholders()
    _table_ready = True
    logger.info("Tabela tabPapelariaModelos verificada/criada")


def _migrate_cabecalho_genero_placeholders() -> None:
    """Atualiza cabeçalhos legados 'Gabinete do Deputado' para placeholders derivados."""
    for col in ("PreviewConfigJson", "CamposJson", "ConfigCampos"):
        if not _table_has_column(col):
            continue
        try:
            db.execute_query(
                f"""
                UPDATE tabPapelariaModelos
                SET {col} = REPLACE({col}, %(_legado)s, %(_novo)s)
                WHERE {col} IS NOT NULL
                  AND {col} LIKE %(_like)s
                """,
                {
                    "_legado": _CABECALHO_LEGADO,
                    "_novo": _CABECALHO_DERIVADO,
                    "_like": f"%{_CABECALHO_LEGADO}%",
                },
            )
        except Exception as exc:
            logger.warning("Migração cabeçalho gênero em %s: %s", col, exc)


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------


def carregar_modelos_da_planilha() -> list[dict[str, Any]]:
    if not SPREADSHEET_PATH.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {SPREADSHEET_PATH}")

    from openpyxl import load_workbook

    wb = load_workbook(SPREADSHEET_PATH, data_only=True)
    ws = wb[ABA_PLANILHA]
    linha_campos = 2
    col_produto = 1
    col_frente_verso = 2
    col_cores = 3
    col_arquivo_modelo = 4
    col_inicio_campos = 5

    modelos: list[dict[str, Any]] = []

    for row in range(3, ws.max_row + 1):
        produto = ws.cell(row=row, column=col_produto).value
        frente_verso = ws.cell(row=row, column=col_frente_verso).value
        cores = ws.cell(row=row, column=col_cores).value
        arquivo_modelo = ws.cell(row=row, column=col_arquivo_modelo).value

        if not arquivo_modelo:
            continue

        arquivo_modelo = str(arquivo_modelo).strip()
        if not arquivo_modelo.lower().endswith(".cdr"):
            arquivo_modelo += ".cdr"

        campos: dict[str, str] = {}
        for col in range(col_inicio_campos, ws.max_column + 1):
            nome_campo_corel = ws.cell(row=linha_campos, column=col).value
            valor_modelo = ws.cell(row=row, column=col).value
            if not nome_campo_corel or valor_modelo is None or str(valor_modelo).strip() == "":
                continue
            campos[str(nome_campo_corel).strip()] = str(valor_modelo)

        if not campos:
            logger.info("Linha %s ignorada (sem campos válidos): %s", row, arquivo_modelo)
            continue

        variaveis: set[str] = set()
        for texto in campos.values():
            variaveis.update(extrair_variaveis(texto))

        modelos.append(
            {
                "linha_planilha": row,
                "produto": str(produto or "").strip(),
                "frente_verso": str(frente_verso or "").strip(),
                "cores": str(cores or "").strip(),
                "arquivo_modelo": arquivo_modelo,
                "campos": campos,
                "variaveis": sorted(variaveis),
            }
        )

    return modelos


def importar_planilha() -> dict[str, int]:
    ensure_papelaria_modelos_table()
    planilha = carregar_modelos_da_planilha()

    criados = 0
    atualizados = 0
    ignorados = 0

    has_legacy = _table_has_column("NomeProduto")

    for idx, modelo in enumerate(planilha):
        ordem = (idx + 1) * 10
        params = {
            "produto": modelo["produto"],
            "frente_verso": modelo["frente_verso"],
            "cores": modelo["cores"],
            "arquivo": modelo["arquivo_modelo"],
            "campos": json.dumps(modelo["campos"], ensure_ascii=False),
            "variaveis": json.dumps(modelo["variaveis"], ensure_ascii=False),
            "ordem": ordem,
        }

        existing = db.execute_query(
            """
            SELECT ID FROM tabPapelariaModelos
            WHERE Produto = %(produto)s
              AND FrenteVerso = %(frente_verso)s
              AND Cores = %(cores)s
              AND ArquivoModelo = %(arquivo)s
            LIMIT 1
            """,
            params,
        )

        if existing:
            modelo_id = existing[0]["ID"]
            if has_legacy:
                db.execute_query(
                    """
                    UPDATE tabPapelariaModelos SET
                        CamposJson = %(campos)s,
                        VariaveisJson = %(variaveis)s,
                        Ativo = 1,
                        Ordem = %(ordem)s,
                        NomeProduto = %(produto)s,
                        NomeArquivo = %(arquivo)s,
                        ConfigCampos = %(campos)s
                    WHERE ID = %(id)s
                    """,
                    {**params, "id": modelo_id},
                )
            else:
                db.execute_query(
                    """
                    UPDATE tabPapelariaModelos SET
                        CamposJson = %(campos)s,
                        VariaveisJson = %(variaveis)s,
                        Ativo = 1,
                        Ordem = %(ordem)s
                    WHERE ID = %(id)s
                    """,
                    {**params, "id": modelo_id},
                )
            atualizados += 1
        else:
            if has_legacy:
                db.execute_query(
                    """
                    INSERT INTO tabPapelariaModelos (
                        Produto, FrenteVerso, Cores, ArquivoModelo,
                        CamposJson, VariaveisJson, Ativo, Ordem,
                        NomeProduto, NomeArquivo, ConfigCampos
                    ) VALUES (
                        %(produto)s, %(frente_verso)s, %(cores)s, %(arquivo)s,
                        %(campos)s, %(variaveis)s, 1, %(ordem)s,
                        %(produto)s, %(arquivo)s, %(campos)s
                    )
                    """,
                    params,
                )
            else:
                db.execute_query(
                    """
                    INSERT INTO tabPapelariaModelos (
                        Produto, FrenteVerso, Cores, ArquivoModelo,
                        CamposJson, VariaveisJson, Ativo, Ordem
                    ) VALUES (
                        %(produto)s, %(frente_verso)s, %(cores)s, %(arquivo)s,
                        %(campos)s, %(variaveis)s, 1, %(ordem)s
                    )
                    """,
                    params,
                )
            criados += 1

    return {"criados": criados, "atualizados": atualizados, "ignorados": ignorados}


def _ensure_modelos_importados_se_vazio() -> None:
    ensure_papelaria_modelos_table()
    rows = db.execute_query("SELECT COUNT(*) AS cnt FROM tabPapelariaModelos WHERE Ativo = 1")
    if rows and rows[0].get("cnt", 0) > 0:
        return
    if not SPREADSHEET_PATH.is_file():
        logger.warning("Planilha papelaria ausente; modelos não importados automaticamente.")
        return
    try:
        result = importar_planilha()
        logger.info("Importação automática papelaria: %s", result)
    except Exception as exc:
        logger.error("Falha na importação automática da planilha papelaria: %s", exc)


# ---------------------------------------------------------------------------
# Leitura DB
# ---------------------------------------------------------------------------


def _json_column_from_row(row: dict[str, Any], column: str) -> Optional[str]:
    if column not in row:
        return None
    raw = row.get(column)
    if raw is None:
        return None
    if isinstance(raw, str):
        return raw.strip() or None
    if isinstance(raw, dict):
        return json.dumps(raw, ensure_ascii=False)
    return None


def _mockup_config_json_from_row(row: dict[str, Any]) -> Optional[str]:
    return _json_column_from_row(row, "MockupConfigJson")


def _preview_config_json_from_row(row: dict[str, Any]) -> Optional[str]:
    return _json_column_from_row(row, "PreviewConfigJson")


def _models_select_columns() -> str:
    cols = [
        "ID",
        "Produto",
        "FrenteVerso",
        "Cores",
        "ArquivoModelo",
        "CamposJson",
        "VariaveisJson",
        "Ativo",
        "Ordem",
    ]
    for legacy in ("NomeProduto", "NomeArquivo", "ConfigCampos"):
        if _table_has_column(legacy):
            cols.append(legacy)
    if _table_has_column("MockupConfigJson"):
        cols.append("MockupConfigJson")
    if _table_has_column("PreviewConfigJson"):
        cols.append("PreviewConfigJson")
    if _table_has_column("NomeExibicao"):
        cols.append("NomeExibicao")
    return ", ".join(cols)


def _build_nome_exibicao_padrao(produto: str, frente: str, cores: str) -> str:
    nome_exibicao = produto
    if frente:
        nome_exibicao += f" — {frente}"
    if cores:
        nome_exibicao += f" — {cores}"
    return nome_exibicao


def _resolve_nome_exibicao(
    row: dict[str, Any],
    produto: str,
    frente: str,
    cores: str,
) -> str:
    if _table_has_column("NomeExibicao"):
        custom = _as_str(row.get("NomeExibicao"))
        if custom:
            return custom
    return _build_nome_exibicao_padrao(produto, frente, cores)


def _row_to_modelo(row: dict[str, Any], include_campos: bool = False) -> dict[str, Any]:
    campos = _parse_json_field(row.get("CamposJson"))
    if not campos and row.get("ConfigCampos") is not None:
        campos = _parse_json_field(row.get("ConfigCampos"))
    if isinstance(campos, list):
        campos = {f.get("key", f.get("nome", "")): f.get("valor", "") for f in campos if isinstance(f, dict)}

    variaveis = _parse_json_field(row.get("VariaveisJson"))
    if isinstance(campos, dict) and campos:
        if include_campos:
            campos, variaveis, _ = normalize_campos_from_inspect(campos)
        else:
            vars_from_text: set[str] = set()
            for texto in campos.values():
                vars_from_text.update(extrair_variaveis(str(texto)))
            if vars_from_text:
                if not variaveis:
                    variaveis = _order_variaveis_list(list(vars_from_text))
            elif not variaveis:
                _, variaveis, _ = normalize_campos_from_inspect(campos)
    if not isinstance(variaveis, list):
        variaveis = []

    produto = (row.get("Produto") or row.get("NomeProduto") or "").strip()
    arquivo = (row.get("ArquivoModelo") or row.get("NomeArquivo") or "").strip()
    frente = (row.get("FrenteVerso") or "").strip()
    cores = (row.get("Cores") or "").strip()

    nome_exibicao = _resolve_nome_exibicao(row, produto, frente, cores)

    mockup_json = _mockup_config_json_from_row(row)
    preview_json = _preview_config_json_from_row(row)

    out: dict[str, Any] = {
        "id": row["ID"],
        "produto": produto,
        "frente_versao": frente,
        "cores": cores,
        "arquivo_modelo": arquivo,
        "variaveis": variaveis,
        "nome_exibicao": nome_exibicao,
        "ordem": row.get("Ordem") or 0,
        "mockup_config_json": mockup_json,
        "tem_mockup_config": bool(mockup_json),
        "preview_config_json": preview_json,
        "tem_preview_config": bool(preview_json),
    }
    if include_campos:
        out["campos"] = campos if isinstance(campos, dict) else {}
        out["campos_formulario"] = _build_campos_formulario_com_preview(
            variaveis,
            preview_json,
            campos if isinstance(campos, dict) else None,
        )
    return out


def _infer_campo_labels(
    campos: Optional[dict[str, str]],
    variaveis: list[str],
    preview_config_json: Optional[str],
) -> dict[str, str]:
    labels: dict[str, str] = {}
    if preview_config_json:
        try:
            from services import papelaria_live_preview_service as preview_svc

            parsed = preview_svc.parse_preview_config_json(preview_config_json)
            raw_labels = parsed.get("campoLabels") if parsed else None
            if isinstance(raw_labels, dict):
                labels = {str(k): str(v) for k, v in raw_labels.items() if str(k).strip()}
        except ValueError:
            pass

    if labels or not campos:
        return labels

    used_slugs: set[str] = set()
    for nome_campo, texto_raw in campos.items():
        texto = str(texto_raw or "").strip()
        if extrair_variaveis(texto):
            continue
        if resolve_canonical_campo_template(nome_campo, texto):
            continue
        slug = slugify_variavel_campo(nome_campo, used_slugs)
        if slug in variaveis:
            labels[slug] = str(nome_campo)
    return labels


def _build_campos_formulario_com_preview(
    variaveis: list[str],
    preview_config_json: Optional[str],
    campos: Optional[dict[str, str]] = None,
) -> list[dict[str, Any]]:
    """Formulário base + campos dinâmicos opcionais da prévia (retrocompatível)."""
    labels = _infer_campo_labels(campos, variaveis, preview_config_json)
    base = build_campos_formulario(variaveis, labels)
    if not preview_config_json:
        return base
    try:
        from services import papelaria_live_preview_service as preview_svc

        parsed = preview_svc.parse_preview_config_json(preview_config_json)
    except ValueError:
        return base
    if not parsed or not parsed.get("habilitarCamposDinamicos"):
        return base
    dinamicos = parsed.get("camposDinamicos") or []
    if not isinstance(dinamicos, list):
        return base
    keys_existentes = {f["key"] for f in base}
    for item in dinamicos:
        if not isinstance(item, dict):
            continue
        variavel = str(item.get("variavel") or item.get("key") or "").strip()
        if not variavel or variavel in keys_existentes or variavel in VARIAVEIS_DERIVADAS:
            continue
        field_type = item.get("tipo") or ("email" if variavel == "email" else "text")
        base.append(
            {
                "key": variavel,
                "label": item.get("label") or ROTULOS_VARIAVEIS.get(variavel, variavel),
                "type": field_type,
                "required": bool(item.get("required", False)),
                "rows": 3 if field_type == "textarea" else None,
                "dinamico": True,
            }
        )
        keys_existentes.add(variavel)
    return base


def list_modelos(
    produto: Optional[str] = None,
    cores: Optional[str] = None,
    frente_verso: Optional[str] = None,
) -> list[dict[str, Any]]:
    _ensure_modelos_importados_se_vazio()
    ensure_papelaria_modelos_table()

    query = f"""
        SELECT {_models_select_columns()}
        FROM tabPapelariaModelos
        WHERE Ativo = 1
    """
    params: dict[str, Any] = {}
    if produto:
        query += " AND (Produto = %(produto)s OR NomeProduto = %(produto)s)"
        params["produto"] = produto.strip()
    if cores:
        query += " AND Cores = %(cores)s"
        params["cores"] = cores.strip()
    if frente_verso:
        query += " AND FrenteVerso = %(frente_verso)s"
        params["frente_verso"] = frente_verso.strip()
    query += " ORDER BY Ordem ASC, Produto ASC, FrenteVerso ASC, Cores ASC"

    rows = db.execute_query(query, params or None)
    return [_row_to_modelo(r) for r in rows]


def get_modelo_by_id(modelo_id: int) -> Optional[dict[str, Any]]:
    ensure_papelaria_modelos_table()
    rows = db.execute_query(
        f"""
        SELECT {_models_select_columns()}
        FROM tabPapelariaModelos
        WHERE ID = %(id)s AND Ativo = 1
        LIMIT 1
        """,
        {"id": modelo_id},
    )
    if not rows:
        return None
    return _attach_preview_faces(_row_to_modelo(rows[0], include_campos=True))


def list_modelos_admin() -> list[dict[str, Any]]:
    """Lista modelos para admin (inclui flag de mockup)."""
    ensure_papelaria_modelos_table()
    rows = db.execute_query(
        f"""
        SELECT {_models_select_columns()}
        FROM tabPapelariaModelos
        WHERE Ativo = 1
        ORDER BY Ordem ASC, Produto ASC, FrenteVerso ASC, Cores ASC
        """
    )
    return [_attach_preview_faces(_row_to_modelo(r)) for r in rows]


def get_modelo_admin_by_id(modelo_id: int) -> Optional[dict[str, Any]]:
    ensure_papelaria_modelos_table()
    rows = db.execute_query(
        f"""
        SELECT {_models_select_columns()}
        FROM tabPapelariaModelos
        WHERE ID = %(id)s AND Ativo = 1
        LIMIT 1
        """,
        {"id": modelo_id},
    )
    if not rows:
        return None
    return _row_to_modelo(rows[0], include_campos=False)


def update_modelo_nome_exibicao(
    modelo_id: int,
    nome_exibicao: Optional[str],
) -> dict[str, Any]:
    """Salva ou limpa o nome amigável do modelo (NULL = nome automático)."""
    ensure_papelaria_modelos_table()
    if not _table_has_column("NomeExibicao"):
        raise RuntimeError("Coluna NomeExibicao não disponível no banco.")

    nome = _as_str(nome_exibicao) or None
    if nome and len(nome) > 255:
        raise ValueError("Nome de exibição deve ter no máximo 255 caracteres.")

    existing = db.execute_query(
        "SELECT ID FROM tabPapelariaModelos WHERE ID = %(id)s AND Ativo = 1 LIMIT 1",
        {"id": modelo_id},
    )
    if not existing:
        raise LookupError("Modelo não encontrado.")

    db.execute_query(
        """
        UPDATE tabPapelariaModelos
        SET NomeExibicao = %(nome)s
        WHERE ID = %(id)s
        """,
        {"id": modelo_id, "nome": nome},
    )

    updated = get_modelo_admin_by_id(modelo_id)
    if not updated:
        raise LookupError("Modelo não encontrado após atualização.")
    return updated


def update_modelo_mockup_config(
    modelo_id: int,
    mockup_config_json: Any,
) -> dict[str, Any]:
    """Salva ou limpa MockupConfigJson do modelo."""
    ensure_papelaria_modelos_table()
    if not _table_has_column("MockupConfigJson"):
        raise RuntimeError("Coluna MockupConfigJson não disponível no banco.")

    from services import papelaria_mockup_config_service as mockup_svc

    normalized = mockup_svc.normalize_mockup_config_json(mockup_config_json)

    existing = db.execute_query(
        "SELECT ID FROM tabPapelariaModelos WHERE ID = %(id)s AND Ativo = 1 LIMIT 1",
        {"id": modelo_id},
    )
    if not existing:
        raise LookupError("Modelo não encontrado.")

    db.execute_query(
        """
        UPDATE tabPapelariaModelos
        SET MockupConfigJson = %(cfg)s
        WHERE ID = %(id)s
        """,
        {"id": modelo_id, "cfg": normalized},
    )

    updated = get_modelo_admin_by_id(modelo_id)
    if not updated:
        raise LookupError("Modelo não encontrado após atualização.")
    return updated


def validate_modelo_mockup_config(mockup_config_json: Any) -> dict[str, Any]:
    from services import papelaria_mockup_config_service as mockup_svc

    parsed = mockup_svc.parse_mockup_config_json(mockup_config_json)
    if parsed is None:
        return {"valid": True, "message": "Configuração vazia (usará preset automático)."}
    mockup_svc.validate_mockup_config_payload(parsed)
    return {"valid": True, "message": "Configuração válida."}


def update_modelo_preview_config(
    modelo_id: int,
    preview_config_json: Any,
) -> dict[str, Any]:
    """Salva ou limpa PreviewConfigJson do modelo (prévia em tempo real DEV)."""
    ensure_papelaria_modelos_table()
    if not _table_has_column("PreviewConfigJson"):
        raise RuntimeError("Coluna PreviewConfigJson não disponível no banco.")

    from services import papelaria_live_preview_service as preview_svc

    normalized = preview_svc.normalize_preview_config_json(preview_config_json)

    existing = db.execute_query(
        "SELECT ID FROM tabPapelariaModelos WHERE ID = %(id)s AND Ativo = 1 LIMIT 1",
        {"id": modelo_id},
    )
    if not existing:
        raise LookupError("Modelo não encontrado.")

    db.execute_query(
        """
        UPDATE tabPapelariaModelos
        SET PreviewConfigJson = %(cfg)s
        WHERE ID = %(id)s
        """,
        {"id": modelo_id, "cfg": normalized},
    )

    updated = get_modelo_admin_by_id(modelo_id)
    if not updated:
        raise LookupError("Modelo não encontrado após atualização.")
    return updated


def delete_modelo_admin(modelo_id: int) -> None:
    """Desativa modelo (soft delete — não aparece mais na lista nem na Papelaria)."""
    ensure_papelaria_modelos_table()
    existing = db.execute_query(
        "SELECT ID FROM tabPapelariaModelos WHERE ID = %(id)s AND Ativo = 1 LIMIT 1",
        {"id": modelo_id},
    )
    if not existing:
        raise LookupError("Modelo não encontrado.")

    db.execute_query(
        """
        UPDATE tabPapelariaModelos
        SET Ativo = 0
        WHERE ID = %(id)s
        """,
        {"id": modelo_id},
    )


def validate_modelo_preview_config(preview_config_json: Any) -> dict[str, Any]:
    from services import papelaria_live_preview_service as preview_svc

    parsed = preview_svc.parse_preview_config_json(preview_config_json)
    if parsed is None:
        return {"valid": True, "message": "Configuração vazia (sem prévia em tempo real)."}
    preview_svc.validate_preview_config_payload(parsed)
    return {"valid": True, "message": "Configuração de prévia válida."}


def _attach_preview_faces(modelo: dict[str, Any]) -> dict[str, Any]:
    preview_json = modelo.get("preview_config_json")
    if not preview_json:
        modelo["preview_faces"] = []
        return modelo
    from services import papelaria_live_preview_service as preview_svc

    try:
        parsed = preview_svc.parse_preview_config_json(preview_json)
        from config_manager import config_manager

        api_prefix = f"{config_manager.get_api_path_prefix()}/papelaria"
        modelo["preview_faces"] = preview_svc.list_preview_faces_public(
            int(modelo["id"]),
            parsed,
            api_prefix=api_prefix,
        )
    except Exception:
        modelo["preview_faces"] = []
    return modelo


def get_modelo_with_preview(modelo_id: int) -> Optional[dict[str, Any]]:
    modelo = get_modelo_by_id(modelo_id)
    if not modelo:
        return None
    return _attach_preview_faces(modelo)


def build_campos_formulario(
    variaveis: list[str],
    labels: Optional[dict[str, str]] = None,
) -> list[dict[str, Any]]:
    ordenadas = _order_variaveis_list(variaveis)
    label_map = labels or {}

    fields: list[dict[str, Any]] = []
    for var in ordenadas:
        if var in VARIAVEIS_DERIVADAS:
            continue
        field_type = "email" if var == "email" else "text"
        if var in ("contato_assessor", "funcao_assessor"):
            field_type = "textarea"
        fields.append(
            {
                "key": var,
                "label": label_map.get(var) or ROTULOS_VARIAVEIS.get(var, var.replace("_", " ")),
                "type": field_type,
                "required": var in ("nome_deputado", "genero", "partido", "uf", "gabinete"),
                "rows": 3 if field_type == "textarea" else None,
            }
        )
    return fields


def build_previa_textual(campos: dict[str, str], dados: dict[str, str]) -> dict[str, str]:
    previa: dict[str, str] = {}
    for nome_campo, texto_modelo in campos.items():
        fv = dividir_frente_verso(texto_modelo)
        if fv:
            previa[nome_campo] = (
                f"[Frente]\n{substituir_variaveis(fv[1], dados)}\n\n"
                f"[Verso]\n{substituir_variaveis(fv[2], dados)}"
            )
        else:
            previa[nome_campo] = substituir_variaveis(texto_modelo, dados) or ""
    return previa


# ---------------------------------------------------------------------------
# Matching OS → modelos
# ---------------------------------------------------------------------------


def _normalizar_cores_os(cores: str) -> str:
    c = (cores or "").strip().lower()
    if "relev" in c or "relevo" in c:
        return "relev seco"
    if "ingl" in c or "port" in c:
        return "port/inglês"
    if "poli" in c or "cor" in c:
        return "policromia"
    return c


def is_relevo_seco_cores(cores: str | None) -> bool:
    """True quando a variante de cores indica relevo seco."""
    return _normalizar_cores_os(cores or "") == "relev seco"


def _normalizar_frente_verso_os(fv: str) -> str:
    """
    Canonicaliza Frente/Verso da OS para matching com modelos.
    tabDetalhesServico.FrenteVerso é boolean (0/1) no MySQL — não texto livre.
    """
    v = (fv or "").strip().lower()
    if not v:
        return ""

    if v in ("1", "true", "yes", "sim", "s", "on"):
        return "frente e verso"
    if v in ("0", "false", "no", "nao", "não", "n", "off"):
        return "somente frente"

    if (
        "frente e verso" in v
        or "frente/verso" in v
        or "f/v" in v
        or re.search(r"\bfv\b", v)
    ):
        return "frente e verso"
    if "somente frente" in v or "só frente" in v or "so frente" in v:
        return "somente frente"
    if v == "frente":
        return "somente frente"
    return ""


def _rotulo_frente_verso_os(fv: Any) -> str:
    """Rótulo legível para exibição na UI."""
    norm = _normalizar_frente_verso_os(_as_str(fv))
    if norm == "frente e verso":
        return "Frente e Verso"
    if norm == "somente frente":
        return "Somente Frente"
    raw = _as_str(fv)
    return raw if raw else ""


def match_modelos_para_os(
    produto_os: str,
    cores_os: str = "",
    frente_verso_os: str = "",
    modelos: Optional[list[dict[str, Any]]] = None,
) -> tuple[list[dict[str, Any]], list[str], Optional[int]]:
    produto_os = _as_str(produto_os)
    cores_os = _as_str(cores_os)
    frente_verso_os = _as_str(frente_verso_os)
    warnings: list[str] = []
    todos = modelos if modelos is not None else list_modelos()
    if not todos:
        warnings.append("Nenhum modelo Corel cadastrado. Importe a planilha de modelos.")
        return [], warnings, None

    prod_norm = normalizar_texto_produto(produto_os)
    if not prod_norm:
        warnings.append("Produto da OS não identificado.")
        return [], warnings, None

    exatos = [m for m in todos if normalizar_texto_produto(m["produto"]) == prod_norm]
    candidatos = exatos

    if not candidatos:
        parciais = [
            m
            for m in todos
            if prod_norm in normalizar_texto_produto(m["produto"])
            or normalizar_texto_produto(m["produto"]) in prod_norm
        ]
        if len(parciais) == 1:
            candidatos = parciais
            warnings.append(
                f"Correspondência parcial de produto: '{parciais[0]['produto']}' para OS '{produto_os}'."
            )
        elif len(parciais) > 1:
            candidatos = parciais
            warnings.append(
                f"Múltiplas correspondências parciais para '{produto_os}'. Selecione o modelo manualmente."
            )
        else:
            warnings.append(f"Nenhum modelo encontrado para o produto '{produto_os}'.")
            return [], warnings, None

    cores_norm = _normalizar_cores_os(cores_os)
    fv_norm = _normalizar_frente_verso_os(frente_verso_os)

    filtrados = candidatos
    if cores_os.strip():
        por_cores = [
            m for m in candidatos if _normalizar_cores_os(m["cores"]) == cores_norm
        ]
        if por_cores:
            filtrados = por_cores
        else:
            warnings.append(
                f"Cores da OS ('{cores_os}') não batem com modelos; exibindo todas as variações do produto."
            )

    if fv_norm:
        por_fv = [
            m
            for m in filtrados
            if _normalizar_frente_verso_os(m["frente_versao"]) == fv_norm
        ]
        if por_fv:
            filtrados = por_fv
        else:
            warnings.append(
                f"Frente/Verso da OS ('{_rotulo_frente_verso_os(frente_verso_os)}') "
                f"não bate com modelos; exibindo demais variações do produto."
            )

    auto_id: Optional[int] = None
    if len(filtrados) == 1:
        auto_id = filtrados[0]["id"]
    elif len(filtrados) > 1:
        warnings.append(
            f"{len(filtrados)} modelos compatíveis. Selecione cores/variação desejada."
        )

    return filtrados, warnings, auto_id


# ---------------------------------------------------------------------------
# Contexto OS + prefill Corel
# ---------------------------------------------------------------------------


def _parse_partido_uf(sigla: str, titulo: str) -> tuple[str, str]:
    raw = (sigla or "").strip()
    if not raw and titulo and "/" in titulo:
        raw = titulo.strip()
    if "/" in raw:
        partes = raw.split("/", 1)
        return partes[0].strip(), partes[1].strip().upper()
    return raw, ""


def build_prefill_corel(ano: int, nro: int) -> dict[str, str]:
    rows = db.execute_query(
        """
        SELECT p.NomeUsuario, p.Titular, p.SiglaOrgao, p.GabSalaUsuario,
               p.RamalUsuario, p.Localizacao, p.Andar,
               d.Titulo, d.TipoPublicacaoLink, d.Cores, d.CoresDescricao,
               d.FrenteVerso
        FROM tabProtocolos p
        LEFT JOIN tabDetalhesServico d
          ON p.NroProtocolo = d.NroProtocoloLinkDet
         AND p.AnoProtocolo = d.AnoProtocoloLinkDet
        WHERE p.NroProtocolo = %(nro)s AND p.AnoProtocolo = %(ano)s
        """,
        {"nro": nro, "ano": ano},
    )
    if not rows:
        return {}

    row = rows[0]
    nome = _as_str(row.get("Titular") or row.get("NomeUsuario"))
    partido, uf = _parse_partido_uf(
        _as_str(row.get("SiglaOrgao")),
        _as_str(row.get("Titulo")),
    )

    gab_raw = _as_str(row.get("GabSalaUsuario"))
    gabinete = normalizar_gabinete(gab_raw) if gab_raw else ""

    anexo = "IV"
    if gab_raw and re.search(r"anexo\s*iv\b", gab_raw, re.I):
        anexo = "IV"

    email = ""
    if nome:
        slug = re.sub(r"[^a-z0-9]", "", nome.split()[0].lower()) if nome.split() else ""
        if slug:
            email = f"dep.{slug}@camara.leg.br"

    dados: dict[str, str] = {
        "nome_deputado": nome,
        "genero": "",
        "partido": partido,
        "uf": uf.upper() if uf else "",
        "anexo": anexo,
        "gabinete": gabinete,
        "email": email,
        "nome_assessor": "",
        "contato_assessor": "",
        "funcao_assessor": "",
    }

    if nome and len(nome) >= 3:
        try:
            from services.deputados_camara_service import (
                aplicar_deputado_em_campos_corel,
                buscar_deputado_por_nome,
            )

            dep = buscar_deputado_por_nome(nome)
            if dep:
                dados = aplicar_deputado_em_campos_corel(dados, dep)
        except Exception as exc:
            logger.warning("Prefill deputado (Corel) via API Câmara falhou: %s", exc)

    if not (dados.get("genero") or "").strip():
        dados["genero"] = "Deputado Federal"

    return dados


def is_papelaria_os_by_product(ano: int, nro: int) -> bool:
    rows = db.execute_query(
        """
        SELECT d.TipoPublicacaoLink
        FROM tabProtocolos p
        LEFT JOIN tabDetalhesServico d
          ON p.NroProtocolo = d.NroProtocoloLinkDet
         AND p.AnoProtocolo = d.AnoProtocoloLinkDet
        WHERE p.NroProtocolo = %(nro)s AND p.AnoProtocolo = %(ano)s
        """,
        {"nro": nro, "ano": ano},
    )
    if not rows:
        return False
    produto = _as_str(rows[0].get("TipoPublicacaoLink"))
    if not produto:
        return False
    
    prod_norm = normalizar_texto_produto(produto)
    if not prod_norm:
        return False
    
    todos = list_modelos()
    return any(normalizar_texto_produto(m["produto"]) == prod_norm for m in todos)


def build_os_context(ano: int, nro: int) -> dict[str, Any]:
    if nro < 5000 and not is_papelaria_os_by_product(ano, nro):
        raise ValueError("OS de papelaria deve ter número >= 5000 ou possuir produto de papelaria configurado.")

    rows = db.execute_query(
        """
        SELECT p.NroProtocolo, p.AnoProtocolo, p.NomeUsuario, p.Titular,
               p.SiglaOrgao, p.GabSalaUsuario, p.RamalUsuario,
               d.Titulo, d.TipoPublicacaoLink, d.Cores, d.CoresDescricao, d.FrenteVerso,
               d.Observ
        FROM tabProtocolos p
        LEFT JOIN tabDetalhesServico d
          ON p.NroProtocolo = d.NroProtocoloLinkDet
         AND p.AnoProtocolo = d.AnoProtocoloLinkDet
        WHERE p.NroProtocolo = %(nro)s AND p.AnoProtocolo = %(ano)s
        """,
        {"nro": nro, "ano": ano},
    )
    if not rows:
        raise LookupError("OS não encontrada.")

    row = rows[0]
    produto = _as_str(row.get("TipoPublicacaoLink"))
    cores = _as_str(row.get("CoresDescricao") or row.get("Cores"))
    frente_verso = _as_str(row.get("FrenteVerso"))

    prefill = build_prefill_corel(ano, nro)
    cliente_alteracao = None
    alt = None
    try:
        from services.proof_papelaria_service import load_cliente_alteracao_prova

        alt = load_cliente_alteracao_prova(ano, nro)
    except Exception as exc:
        logger.warning("Falha ao carregar alteração de prova do cliente: %s", exc)
        alt = None

    if alt:
        campos_alt = alt.get("campos") or {}
        job_dados_base: dict[str, str] = {}
        job_id_raw = alt.get("job_id")
        if job_id_raw:
            try:
                from services.papelaria_jobs_service import get_job

                job = get_job(int(job_id_raw))
                if job and isinstance(job.get("dados"), dict):
                    job_dados_base = {
                        str(k): str(v) for k, v in job["dados"].items()
                    }
            except Exception as exc:
                logger.warning(
                    "Falha ao carregar job da alteração do cliente: %s", exc
                )
        if campos_alt:
            if job_dados_base:
                prefill = {**job_dados_base, **campos_alt}
            else:
                prefill = {**prefill, **campos_alt}
        elif job_dados_base:
            prefill = {**prefill, **job_dados_base}

        obs_cli = (alt.get("observacao_cliente") or "").strip()
        campos_alterados: list[dict[str, str]] = []
        if campos_alt:
            for key, depois in campos_alt.items():
                antes = job_dados_base.get(str(key), "").strip()
                depois_s = str(depois).strip()
                if antes != depois_s:
                    campos_alterados.append(
                        {
                            "key": str(key),
                            "antes": antes,
                            "depois": depois_s,
                        }
                    )
        if obs_cli or campos_alt:
            cliente_alteracao = {
                "observacao": obs_cli,
                "versao": alt.get("versao") or "",
                "modelo_id": alt.get("modelo_id"),
                "job_id": alt.get("job_id"),
                "campos": campos_alt if campos_alt else None,
                "campos_alterados": campos_alterados or None,
            }

    todos_modelos = list_modelos()
    modelos, warnings, auto_id = match_modelos_para_os(
        produto, cores, frente_verso, todos_modelos
    )

    if alt:
        try:
            alt_modelo_id = int(alt.get("modelo_id") or 0)
        except (TypeError, ValueError):
            alt_modelo_id = 0
        if alt_modelo_id and any(m.get("id") == alt_modelo_id for m in modelos):
            auto_id = alt_modelo_id

    partido, uf = _parse_partido_uf(
        _as_str(row.get("SiglaOrgao")),
        _as_str(row.get("Titulo")),
    )

    return {
        "os": {
            "nro": nro,
            "ano": ano,
            "produto": produto,
            "titulo": _as_str(row.get("Titulo")),
            "solicitante": _as_str(row.get("NomeUsuario")),
            "titular": _as_str(row.get("Titular")),
            "partido": partido,
            "uf": uf,
            "gabinete": _as_str(row.get("GabSalaUsuario")),
            "cores": cores,
            "frente_verso": _rotulo_frente_verso_os(frente_verso),
            "observ": _as_str(row.get("Observ")),
        },
        "prefill": prefill,
        "modelos": modelos,
        "autoSelectedModelId": auto_id,
        "warnings": warnings,
        "cliente_alteracao": cliente_alteracao,
    }
