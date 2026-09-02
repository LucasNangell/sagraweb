from __future__ import annotations

from pathlib import Path

from services.sequencia_producao_models import SequenciaDraft


def write_sequencia_docx(draft: SequenciaDraft, target_path: str | Path) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt
    except ImportError as exc:
        raise RuntimeError(
            "Dependência python-docx não instalada. Instale as dependências do backend."
        ) from exc

    path = Path(target_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SEQUÊNCIA DE PRODUÇÃO")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        f"OS {draft.os.numero}/{draft.os.ano}"
        + (f" - {draft.os.titulo}" if draft.os.titulo else "")
    )
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(11)

    doc.add_paragraph()

    for section_draft in draft.sections:
        for block in section_draft.blocks:
            if not block.title.strip() and not any(line.strip() for line in block.lines):
                continue
            heading = doc.add_paragraph()
            heading_run = heading.add_run(block.title.strip())
            heading_run.bold = True
            heading_run.font.name = "Arial"
            heading_run.font.size = Pt(12)
            for line in block.lines:
                if not line.strip():
                    continue
                paragraph = doc.add_paragraph(style=None)
                paragraph.paragraph_format.left_indent = Cm(0.35)
                paragraph.paragraph_format.space_after = Pt(2)
                run = paragraph.add_run(line.strip())
                run.font.name = "Arial"
                run.font.size = Pt(11)
            doc.add_paragraph()

    doc.save(path)
    return path
