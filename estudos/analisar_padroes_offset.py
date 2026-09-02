from __future__ import annotations

import csv
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from docx import Document


BASE_DIR = Path(__file__).resolve().parent
XLSX_PATH = BASE_DIR / "relatorio trabalhos offset.xlsx"
CSV_OUT = BASE_DIR / "analise_offset_trabalhos.csv"
JSON_OUT = BASE_DIR / "analise_offset_padroes.json"
REPORT_OUT = BASE_DIR / "relatorio_padroes_offset.md"


def norm(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("_x000D_", "\n").replace("\r", "\n")
    return "\n".join(line.strip() for line in text.splitlines() if line.strip())


def key(text: object) -> str:
    value = norm(text).lower()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.replace("º", "o").replace("°", "o")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def split_tokens(text: object) -> list[str]:
    clean = norm(text)
    parts: list[str] = []
    for line in clean.splitlines():
        line = re.sub(r"^[\-\u2022]+\s*", "", line.strip())
        if line:
            parts.append(line)
    return parts


def read_jobs() -> list[dict[str, object]]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [norm(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    jobs: list[dict[str, object]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in row):
            continue
        item = dict(zip(headers, row))
        item["Numero OS"] = int(item.get("Número OS") or item.get("Numero OS") or 0)
        item["Ano"] = int(item.get("Ano") or 0)
        for name in headers:
            if name not in {"Número OS", "Ano", "Tiragem Final", "Páginas", "Frente/Verso"}:
                item[name] = norm(item.get(name))
        jobs.append(item)
    return jobs


def find_doc(job: dict[str, object]) -> tuple[Path | None, str]:
    year = str(job["Ano"])
    os_num = f"{int(job['Numero OS']):05d}"
    folder = BASE_DIR / year
    expected = norm(job.get("Arquivo"))
    if expected:
        candidate = folder / expected
        if candidate.exists():
            return candidate, "arquivo_coluna"
    matches = sorted(folder.glob(f"{os_num}_*.docx"))
    if len(matches) == 1:
        return matches[0], "prefixo_os"
    if matches:
        plain = [p for p in matches if "komori" not in key(p.name) and "capa" not in key(p.name)]
        return (plain[0] if plain else matches[0]), "prefixo_os_multiplos"
    return None, "nao_encontrado"


STAGE_ALIASES = [
    ("Formato de Impressao", [r"^formato de impressao$"]),
    ("Laminacao", [r"^laminacao$", r"laminacao"]),
    ("Hot Stamping", [r"^hot stamping$", r"hot stamping"]),
    ("Corte", [r"^corte$", r"^cortes?$"]),
    ("Dobra", [r"^dobra", r"dobradeira"]),
    ("Alceamento", [r"^alceamento"]),
    ("Encadernacao", [r"^encadernacao", r"canoa", r"grampos?$"]),
    ("Capeadeira", [r"^capeadeira$", r"capa dura"]),
    ("Wire-O", [r"wire o", r"wire"]),
    ("Espiral", [r"espiral"]),
    ("Furo", [r"^furo", r"furacao", r"furos?"]),
    ("Serrilha/Picote", [r"serrilh", r"picot"]),
    ("Refile Final", [r"^refile final", r"guilhotina$"]),
]


def heading_for(text: str) -> str | None:
    k = key(text)
    if k in {"impressao", "acabamento"}:
        return None
    for stage, patterns in STAGE_ALIASES:
        if any(re.search(pattern, k) for pattern in patterns):
            return stage
    return None


def extract_doc(doc_path: Path) -> dict[str, object]:
    doc = Document(str(doc_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    stages: dict[str, list[str]] = {}
    current: str | None = None
    for text in paragraphs:
        stage = heading_for(text)
        if stage:
            current = stage
            stages.setdefault(stage, [])
            continue
        if current:
            stages.setdefault(current, []).append(text)

    all_text = "\n".join(paragraphs)
    k = key(all_text)
    flags = {
        "tem_laminacao": "laminacao" in k,
        "tem_dobra": "dobra" in k or "dobradeira" in k,
        "tem_alceamento": "alceamento" in k,
        "tem_canoa": "canoa" in k or "grampos" in k,
        "tem_wire_o": "wire" in k,
        "tem_espiral": "espiral" in k,
        "tem_capeadeira": "capeadeira" in k,
        "tem_hot_stamping": "hot stamping" in k,
        "tem_serrilha_picote": "serrilh" in k or "picot" in k,
        "tem_refile": "refile" in k or "guilhotina" in k,
    }
    return {
        "paragraphs": paragraphs,
        "stages": {name: values for name, values in stages.items()},
        "stage_order": list(stages.keys()),
        "flags": flags,
    }


def format_counter(counter: Counter, limit: int = 12) -> list[dict[str, object]]:
    return [{"valor": k, "qtde": v} for k, v in counter.most_common(limit)]


def pct(part: int, total: int) -> str:
    return f"{part / total:.0%}" if total else "0%"


def main() -> None:
    jobs = read_jobs()
    rows: list[dict[str, object]] = []
    counters: dict[str, Counter] = defaultdict(Counter)
    stage_by_product: dict[str, Counter] = defaultdict(Counter)
    fold_signatures: Counter[str] = Counter()
    print_formats: Counter[str] = Counter()
    missing_docs: list[dict[str, object]] = []

    for job in jobs:
        doc_path, match_kind = find_doc(job)
        extracted = None
        error = ""
        if doc_path is None:
            missing_docs.append(job)
        else:
            try:
                extracted = extract_doc(doc_path)
            except Exception as exc:  # keep batch processing resilient
                error = f"{type(exc).__name__}: {exc}"

        produto = norm(job.get("Produto (Tipo)"))
        formato = norm(job.get("Formato"))
        maquina = norm(job.get("Máquina"))
        acabamento_tokens = split_tokens(job.get("Acabamento"))
        cores = norm(job.get("Cores"))
        papel = norm(job.get("Papel (Descrição)"))
        paginas = int(job.get("Páginas") or 0)
        fv = int(job.get("Frente/Verso") or 0)

        for field, value in [
            ("produto", produto),
            ("formato", formato),
            ("maquina", maquina),
            ("cores", cores),
            ("papel_tipo", norm(job.get("Papel (Tipo)"))),
        ]:
            counters[field][value or "(vazio)"] += 1
        for token in acabamento_tokens:
            counters["acabamento"][token] += 1

        stage_order = extracted["stage_order"] if extracted else []
        flags = extracted["flags"] if extracted else {}
        stages = extracted["stages"] if extracted else {}

        for stage in stage_order:
            stage_by_product[produto][stage] += 1
        for line in stages.get("Dobra", []) if stages else []:
            if re.search(r"\b(4|8|12|16|24|32)\s*p(?:a?g|gs?)\b", key(line)):
                fold_signatures[line] += 1
        for line in stages.get("Formato de Impressao", []) if stages else []:
            if re.search(r"\d+\s*x\s*\d+", key(line)):
                print_formats[line] += 1

        row = {
            "os": f"{int(job['Numero OS']):05d}/{int(job['Ano'])}",
            "produto": produto,
            "formato": formato,
            "maquina": maquina,
            "tiragem": int(job.get("Tiragem Final") or 0),
            "paginas": paginas,
            "frente_verso": fv,
            "papel": papel.replace("\n", " / "),
            "cores": cores,
            "cores_descricao": norm(job.get("Cores (Descrição)")).replace("\n", " / "),
            "acabamento_planilha": " | ".join(acabamento_tokens),
            "arquivo_docx": str(doc_path) if doc_path else "",
            "match_docx": match_kind,
            "erro_docx": error,
            "etapas_docx": " > ".join(stage_order),
            **{name: bool(flags.get(name)) for name in [
                "tem_laminacao",
                "tem_dobra",
                "tem_alceamento",
                "tem_canoa",
                "tem_wire_o",
                "tem_espiral",
                "tem_capeadeira",
                "tem_hot_stamping",
                "tem_serrilha_picote",
                "tem_refile",
            ]},
            "formato_impressao_docx": " | ".join(stages.get("Formato de Impressao", [])) if stages else "",
            "dobra_docx": " | ".join(stages.get("Dobra", [])) if stages else "",
            "corte_docx": " | ".join(stages.get("Corte", [])) if stages else "",
            "alceamento_docx": " | ".join(stages.get("Alceamento", [])) if stages else "",
            "encadernacao_docx": " | ".join(stages.get("Encadernacao", [])) if stages else "",
            "refile_docx": " | ".join(stages.get("Refile Final", [])) if stages else "",
            "observacoes": norm(job.get("Observações")).replace("\n", " / "),
        }
        rows.append(row)

    fieldnames = list(rows[0].keys()) if rows else []
    with CSV_OUT.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    total = len(rows)
    matched = sum(1 for row in rows if row["arquivo_docx"])
    flag_counts = {
        flag: sum(1 for row in rows if row.get(flag))
        for flag in [
            "tem_laminacao",
            "tem_dobra",
            "tem_alceamento",
            "tem_canoa",
            "tem_wire_o",
            "tem_espiral",
            "tem_capeadeira",
            "tem_hot_stamping",
            "tem_serrilha_picote",
            "tem_refile",
        ]
    }

    by_product_summary = {}
    for product in counters["produto"]:
        product_rows = [row for row in rows if row["produto"] == product]
        by_product_summary[product] = {
            "qtde": len(product_rows),
            "formatos": format_counter(Counter(row["formato"] for row in product_rows), 8),
            "maquinas": format_counter(Counter(row["maquina"] for row in product_rows), 8),
            "paginas": {
                "min": min((row["paginas"] for row in product_rows), default=0),
                "max": max((row["paginas"] for row in product_rows), default=0),
                "valores_mais_comuns": format_counter(Counter(row["paginas"] for row in product_rows), 8),
            },
            "etapas": format_counter(stage_by_product[product], 12),
            "etapas_flags": {
                flag: sum(1 for row in product_rows if row.get(flag))
                for flag in flag_counts
            },
            "acabamentos_planilha": format_counter(Counter(
                token
                for row in product_rows
                for token in str(row["acabamento_planilha"]).split(" | ")
                if token
            ), 12),
        }

    summary = {
        "fontes": {
            "xlsx": str(XLSX_PATH),
            "total_planilha": total,
            "docx_encontrados": matched,
            "docx_nao_encontrados": len(missing_docs),
        },
        "frequencias": {
            "produtos": format_counter(counters["produto"], 20),
            "formatos": format_counter(counters["formato"], 20),
            "maquinas": format_counter(counters["maquina"], 20),
            "cores": format_counter(counters["cores"], 20),
            "acabamentos": format_counter(counters["acabamento"], 30),
            "etapas_docx": {k: {"qtde": v, "percentual": pct(v, total)} for k, v in flag_counts.items()},
            "formatos_impressao_docx": format_counter(print_formats, 30),
            "dobras_docx": format_counter(fold_signatures, 30),
        },
        "por_produto": by_product_summary,
        "linhas": rows,
    }
    JSON_OUT.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    write_report(summary)
    print(f"Linhas analisadas: {total}")
    print(f"DOCX encontrados: {matched}")
    print(f"CSV: {CSV_OUT}")
    print(f"JSON: {JSON_OUT}")
    print(f"Relatorio: {REPORT_OUT}")


def md_table(rows: list[dict[str, object]], columns: list[tuple[str, str]], limit: int | None = None) -> str:
    selected = rows[:limit] if limit else rows
    header = "| " + " | ".join(label for _, label in columns) + " |"
    sep = "| " + " | ".join("---" for _ in columns) + " |"
    body = []
    for row in selected:
        values = []
        for key_name, _ in columns:
            value = str(row.get(key_name, ""))
            value = value.replace("|", "\\|").replace("\n", "<br>")
            values.append(value)
        body.append("| " + " | ".join(values) + " |")
    return "\n".join([header, sep, *body])


def write_report(summary: dict[str, object]) -> None:
    fontes = summary["fontes"]
    freqs = summary["frequencias"]
    produtos = freqs["produtos"]
    etapas = freqs["etapas_docx"]
    por_produto = summary["por_produto"]

    canoa_terms = ["Dobra", "Grampo Canoa", "Refile"]
    linhas = summary["linhas"]
    divergencias = [
        row for row in linhas
        if ("dobra" in key(row["acabamento_planilha"]) and not row["tem_dobra"])
        or ("grampo" in key(row["acabamento_planilha"]) and not row["tem_canoa"])
        or ("wire" in key(row["acabamento_planilha"]) and not row["tem_wire_o"])
        or ("lamina" in key(row["acabamento_planilha"]) and "sem laminacao" not in key(row["acabamento_planilha"]) and not row["tem_laminacao"])
    ]
    sem_dobra = [row for row in linhas if not row["tem_dobra"]]
    sem_alceamento = [row for row in linhas if not row["tem_alceamento"]]
    wire_o = [row for row in linhas if row["tem_wire_o"] or "wire" in key(row["acabamento_planilha"])]

    lines = [
        "# Relatório de padrões - trabalhos offset",
        "",
        f"Fonte principal: `{Path(fontes['xlsx']).name}`.",
        f"Base analisada: {fontes['total_planilha']} OS da planilha; {fontes['docx_encontrados']} `.docx` localizados por ano e número da OS.",
        "",
        "## Leitura geral",
        "",
        "A sequência de produção nos `.docx` é predominantemente textual e segue um modelo bastante estável:",
        "",
        "1. Título com OS e ano.",
        "2. Impressão.",
        "3. Formato de impressão por componente do produto, normalmente `Capa`, `Miolo` e, quando necessário, encartes/divisórias/miolo colorido.",
        "4. Acabamento, com etapas condicionais como laminação, corte, dobra, alceamento, encadernação, capeadeira, Wire-O/espiral, serrilha/picote e refile final.",
        "",
        "## Distribuição da amostra",
        "",
        md_table(produtos, [("valor", "Produto"), ("qtde", "Qtde")], 20),
        "",
        md_table(freqs["formatos"], [("valor", "Formato"), ("qtde", "Qtde")], 15),
        "",
        md_table(freqs["maquinas"], [("valor", "Máquina"), ("qtde", "Qtde")], 15),
        "",
        "## Etapas encontradas nos DOCX",
        "",
        md_table([
            {"etapa": name.replace("tem_", "").replace("_", " "), **data}
            for name, data in etapas.items()
        ], [("etapa", "Etapa"), ("qtde", "Qtde"), ("percentual", "% da base")]),
        "",
        "## Regras candidatas para automação",
        "",
        "### 1. Produto define o esqueleto da sequência",
        "",
        "Use `Produto (Tipo)` como primeira chave de decisão. Ele determina se a OS tende a virar uma sequência de canoa, livro colado/capeado, lâminas soltas, Wire-O ou item de acabamento simples.",
        "",
        "Padrões por produto:",
        "",
    ]

    for product, data in sorted(por_produto.items(), key=lambda item: item[1]["qtde"], reverse=True):
        etapa_flags = data["etapas_flags"]
        etapas_txt = ", ".join(
            f"{label} ({etapa_flags[flag]})"
            for flag, label in [
                ("tem_laminacao", "laminação"),
                ("tem_dobra", "dobra"),
                ("tem_alceamento", "alceamento"),
                ("tem_canoa", "canoa"),
                ("tem_wire_o", "Wire-O"),
                ("tem_capeadeira", "capeadeira"),
                ("tem_refile", "refile"),
            ]
            if etapa_flags[flag]
        )
        formatos_txt = ", ".join(f"{x['valor']} ({x['qtde']})" for x in data["formatos"][:5])
        acab_txt = ", ".join(f"{x['valor']} ({x['qtde']})" for x in data["acabamentos_planilha"][:6])
        lines.extend([
            f"- **{product}** ({data['qtde']} OS): formatos principais: {formatos_txt or 'sem padrão dominante'}; etapas DOCX: {etapas_txt or 'sem extração'}; acabamentos informados: {acab_txt or 'sem acabamento informado'}.",
        ])

    lines.extend([
        "",
        "### 2. Acabamento da planilha liga e parametriza as etapas",
        "",
        "A coluna `Acabamento` deve ser normalizada para uma lista de operações. Termos equivalentes e grafias com erro devem cair no mesmo código interno.",
        "",
        md_table(freqs["acabamentos"], [("valor", "Acabamento informado"), ("qtde", "Qtde")], 30),
        "",
        "Normalizações recomendadas:",
        "",
        "- `Grampo canoa`, `Grampo Canoa`, `Encadernação (Canoa)` -> `ENCADERNACAO_CANOA`.",
        "- `Granpo` -> tratar como `Grampo`.",
        "- `Laminação Brilho/Fosca`, `Laminaçao`, `Laminação na Capa` -> `LAMINACAO`, com atributo `tipo` e `aplicacao`.",
        "- `Wire-O`, `Wire-ô`, `Wire` -> `WIRE_O`.",
        "- `Serrilha`, `picote`, `serilha` -> `SERRILHA_PICOTE`.",
        "- `Refile`, `Refile Final`, `Guilhotina` -> `REFILE_FINAL` quando estiver no fechamento da sequência.",
        "",
        "### 3. Formato de impressão deve ser gerado por componente",
        "",
        "Nos `.docx`, o formato de impressão não é um único campo da OS: ele aparece por componente. Exemplos recorrentes: `Capa`, `Miolo`, `Miolo PB`, `Miolo Colorido`, `Divisória`, `Encarte`.",
        "",
        "Formatos de impressão mais frequentes extraídos dos documentos:",
        "",
        md_table(freqs["formatos_impressao_docx"], [("valor", "Linha do DOCX"), ("qtde", "Qtde")], 20),
        "",
        "Regra de implementação: montar uma lista de componentes a partir de `Papel (Descrição)`, `Cores (Descrição)` e `Produto (Tipo)`, então aplicar uma tabela de imposição/formato por máquina, formato final e componente.",
        "",
        "### 4. Dobra deriva principalmente de páginas por caderno",
        "",
        "As linhas de dobra mostram assinaturas como `16 pgs`, `8 pgs` e `4 pgs`, com instruções de dobra associadas. Para livros e livretos, esta é uma área com bom potencial de automação.",
        "",
        md_table(freqs["dobras_docx"], [("valor", "Linha de dobra"), ("qtde", "Qtde")], 20),
        "",
        "Regras candidatas:",
        "",
        "- Caderno de `16 pgs`: geralmente três dobras, alternando vertical/cabeça/vertical.",
        "- Caderno de `8 pgs`: geralmente duas dobras, cabeça e vertical.",
        "- Caderno de `4 pgs`: geralmente uma dobra vertical.",
        "- Canoa com capa: separar `Capa - 4 pgs`; miolo recebe cadernos conforme total de páginas restante.",
        "- Livro colado: `Capa` deve ir separada para cola/capeadeira; miolo vira cadernos empilhados.",
        "",
        "### 5. Alceamento e encadernação dependem do tipo de acabamento",
        "",
        "O alceamento aparece quando há cadernos, lâminas ou conjunto capa+miolo a ordenar. O texto costuma indicar o recurso, como `AmGraphics` ou `Manual`, e depois a composição do conjunto.",
        "",
        "Regras candidatas:",
        "",
        "- `Grampo canoa`: alceamento de cadernos empilhados e encadernação canoa com 2 grampos, salvo exceção descrita.",
        "- `Livro` com lombada/orelha/laminação de capa: laminação da capa, corte/separação de capa e miolo, dobra do miolo, alceamento manual, capeadeira e refile.",
        "- `Wire-O`: tratar como encadernação própria, normalmente com furação/serrilha/picote quando informado.",
        "- `Lâminas soltas`: evitar gerar dobra/encadernação se o docx ou acabamento indicar lâminas soltas; usar corte, ordenação e refile.",
        "",
        "## Proposta de motor de regras",
        "",
        "Entrada mínima recomendada a partir da OS:",
        "",
        "- número, ano, produto, formato final, máquina, tiragem, páginas, frente/verso;",
        "- componentes de papel: capa, miolo, folhas especiais, divisórias, encartes;",
        "- mapa de cores por componente: policromia, PB, páginas coloridas específicas;",
        "- operações de acabamento normalizadas;",
        "- observações com prioridade de exceção manual.",
        "",
        "Saída sugerida:",
        "",
        "1. `Impressão` com linhas de formato por componente.",
        "2. `Acabamento` com etapas ordenadas.",
        "3. Para cada etapa: recurso/equipamento, parâmetros e instruções de exceção.",
        "",
        "Ordem-base das etapas:",
        "",
        "`Impressão` -> `Laminação/Hot Stamping` -> `Corte` -> `Dobra` -> `Serrilha/Picote/Furo` -> `Alceamento` -> `Encadernação/Capeadeira/Wire-O/Espiral` -> `Refile Final`.",
        "",
        "## Pontos que precisam virar tabela de parametrização",
        "",
        "- Formato de impressão por `máquina + formato final + componente + papel`.",
        "- Capacidade de caderno por produto/formato: 16, 8, 4 páginas e exceções.",
        "- Recurso padrão de alceamento: AmGraphics versus manual.",
        "- Regras de lombada por páginas/papel, quando a OS já informa ou quando deve calcular.",
        "- Aplicação de laminação: capa frente, capa F/V, folhas especiais, tipo brilho/fosca.",
        "- Vocabulário controlado de acabamento para absorver variações de digitação.",
        "",
        "## Divergências e cuidados",
        "",
        f"Foram detectadas {len(divergencias)} possíveis divergências entre acabamento informado na planilha e etapa reconhecida no `.docx`. Isso não significa erro de produção: algumas etapas podem estar descritas com outra palavra ou embutidas em observações.",
        "",
        "Amostra de divergências para revisão:",
        "",
        md_table(divergencias[:20], [
            ("os", "OS"),
            ("produto", "Produto"),
            ("acabamento_planilha", "Acabamento planilha"),
            ("etapas_docx", "Etapas DOCX"),
            ("arquivo_docx", "DOCX"),
        ]),
        "",
        "Casos que merecem regra de exceção:",
        "",
        f"- Sem etapa de dobra reconhecida: {len(sem_dobra)} OS. Em geral são livros com lâminas soltas, folha especial e Wire-O, não cadernos dobrados.",
        f"- Sem etapa de alceamento reconhecida: {len(sem_alceamento)} OS. Conferir antes de usar alceamento como obrigatório para todo livro/livreto.",
        f"- Com Wire-O ou menção a Wire-O: {len(wire_o)} OS. Esse grupo deve usar regra própria e não a regra canoa, mesmo quando o `.docx` traz texto herdado como `Encadernação (Canoa)`.",
        "",
        "Amostra dos casos sem dobra:",
        "",
        md_table(sem_dobra[:10], [
            ("os", "OS"),
            ("produto", "Produto"),
            ("acabamento_planilha", "Acabamento planilha"),
            ("etapas_docx", "Etapas DOCX"),
        ]),
        "",
        "Amostra dos casos com Wire-O:",
        "",
        md_table(wire_o[:10], [
            ("os", "OS"),
            ("produto", "Produto"),
            ("acabamento_planilha", "Acabamento planilha"),
            ("etapas_docx", "Etapas DOCX"),
        ]),
        "",
        "## Arquivos gerados para auditoria",
        "",
        f"- `{CSV_OUT.name}`: uma linha por OS, com campos da planilha, etapas reconhecidas no DOCX e trechos extraídos.",
        f"- `{JSON_OUT.name}`: resumo estruturado com frequências e agregações por produto.",
        f"- `{REPORT_OUT.name}`: este relatório.",
        "",
        "## Recomendações de implementação",
        "",
        "1. Começar pelo gerador de sequência para `Livreto c/ Capa` e `Revista c/ Capa`, porque o fluxo canoa é o mais repetitivo.",
        "2. Em seguida implementar `Livro`, separando capa, miolo, laminação, dobra de cadernos, alceamento, capeadeira e refile.",
        "3. Criar uma camada de normalização de texto antes das regras; sem isso, grafias como `Granpo`, `Laminaçao`, `Wire-ô` e variações de caixa vão gerar exceções falsas.",
        "4. Manter as observações da OS como bloco de exceções manuais no final de cada etapa afetada.",
        "5. Registrar confiança por etapa: `automática`, `automática com observação`, `requer conferência`.",
    ])

    REPORT_OUT.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
